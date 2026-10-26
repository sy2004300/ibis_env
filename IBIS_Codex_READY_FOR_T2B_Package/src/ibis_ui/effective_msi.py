"""One effective MSI snapshot, used by both generation and inspection.

Raw inputs remain read-only. Generated workbooks are immutable, content-addressed
snapshots; manual workbook editing is detected and never silently accepted.
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import shutil
import tempfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from ._compat37 import unlink_missing_ok
from ibis_auto.msi import _columns, discover
from ibis_auto.values import parse_logic
from ibis_auto.ooxml import read_workbook
from .config_adapter import payload_to_core_config
from .errors import ProjectValidationError
from .msi_master import file_sha256, _safe_sheet_name


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def source_payload(payload, workflow):
    """Recover raw mappings from legacy 0.4.0/1 manually merged projects."""
    result = copy.deepcopy(payload)
    state = workflow.get('msi_master', {})
    if state.get('source_root') and any(Path(row['file']).name == 'MSI_Master.xlsx' for row in result['msi_mappings']):
        result['roots']['msi'] = state['source_root']
        by_sheet = {row['master_sheet']: row for row in state.get('legacy_sources', state.get('sources', []))}
        for row in result['msi_mappings']:
            if Path(row['file']).name == 'MSI_Master.xlsx':
                original = by_sheet.get(row['sheet'])
                if not original:
                    raise ProjectValidationError('MSI_TRACE_MISSING', f"模块 {row['module']} 的旧 Master 缺少原始来源记录，请重新指定映射")
                row.update(file=original['source_workbook'], sheet=original['source_sheet'])
    return result


def candidates(payload):
    root = Path(payload['roots']['msi'])
    files = set(root.rglob('*.xlsx')) if root.is_dir() else set()
    files.update(root / row['file'] for row in payload['msi_mappings'])
    result = []
    for file in sorted(files):
        if file.name.startswith('~$') or file.name == 'MSI_Master.xlsx':
            continue
        try:
            for sheet, rows in read_workbook(file).items():
                if rows:
                    result.append({'file': str(file.resolve()), 'sheet': sheet, 'columns': rows[0]})
        except Exception:
            continue  # Explicit selections receive a located error in inspect().
    return result


@dataclass
class EffectiveMSI:
    payload: dict
    mappings: list
    signals: dict
    sheets: dict
    issues: list
    fingerprint: str

    @property
    def errors(self):
        return [item for item in self.issues if item['status'] == 'Error']

    def require_valid(self):
        if self.errors:
            raise ProjectValidationError('MSI_MASTER_INVALID', '；'.join(f"{x['module']}/{x.get('sheet', '')}/{x.get('signal', '')}：{x['message']}" for x in self.errors))


def inspect(payload, workflow):
    raw = source_payload(payload, workflow)
    available = candidates(raw)
    mappings, signals, sheets, issues = [], {}, {}, []
    old = {row['module']: row for row in workflow.get('msi_master', {}).get('sources', []) if 'module' in row}
    used = set()
    for module_row in raw['modules']:
        module = module_row['name']
        sheet = ''
        signal_context = ''
        try:
            explicit = next((row for row in raw['msi_mappings'] if row['module'].casefold() == module.casefold()), None)
            if explicit:
                path = (Path(raw['roots']['msi']) / explicit['file']).resolve()
                sheet, column = explicit['sheet'], explicit['column']
            else:
                norm = lambda s: ''.join(c for c in s.casefold() if c.isalnum())
                matches = [row for row in available if norm(module) in {norm(Path(row['file']).stem), norm(row['sheet'])}]
                if len(matches) != 1:
                    raise ValueError(f'来源候选为 {len(matches)} 个，请在第 10 节选择工作簿和工作表')
                item = matches[0]
                path, sheet = Path(item['file']), item['sheet']
                column = item['columns'][_columns(item['columns'])['signal_name']]
            before = file_sha256(path)
            rows = read_workbook(path).get(sheet)
            if not rows:
                raise ValueError('工作表不存在或为空')
            columns = _columns(rows[0], column)
            names = [row[columns['signal_name']].strip() for row in rows[1:] if len(row) > columns['signal_name'] and row[columns['signal_name']].strip()]
            counts = Counter(name.casefold() for name in names)
            duplicates = sorted({name for name in names if counts[name.casefold()] > 1})
            if duplicates:
                signal_context = ', '.join(duplicates)
                raise ValueError(f'信号名称重复：{signal_context}')
            for row in rows[1:]:
                if len(row) <= columns['signal_name'] or not row[columns['signal_name']].strip():
                    continue
                signal_context = row[columns['signal_name']]
                width = row[columns['width']] if len(row) > columns['width'] else ''
                try:
                    valid_width = bool(width) and float(width).is_integer() and float(width) > 0
                except (ValueError, TypeError):
                    valid_width = False
                if not valid_width:
                    raise ValueError(f"信号 {row[columns['signal_name']]} 的 WIDTH 无效：{width}")
            signal_context = ''
            config = payload_to_core_config(raw)
            config.msi_mappings[module] = (str(path), sheet, column)
            config.attribute_overrides = []
            original = discover(config, module)
            attrs = [row for row in raw['attribute_overrides'] if row['module'].casefold() == module.casefold()]
            if len({row['signal'].casefold() for row in attrs}) != len(attrs):
                signal_context = ', '.join(row['signal'] for row in attrs)
                raise ValueError(f'供电域属性覆盖重复或冲突：{signal_context}')
            config.attribute_overrides = attrs
            effective = discover(config, module)
            by_name = {s.name.casefold(): s for s in effective}
            for override in raw['overrides']:
                signal_context = override['signal'] if override['module'].casefold() == module.casefold() else ''
                if override['module'].casefold() == module.casefold() and override['signal'].casefold() not in by_name:
                    raise ValueError(f"信号值覆盖引用不存在的信号：{override['signal']}")
                if override['module'].casefold() == module.casefold():
                    parse_logic(override['value'], module, override['signal'], by_name[override['signal'].casefold()].width)
            signal_context = ''
            master_sheet = _safe_sheet_name(module, used)
            adjusted = copy.deepcopy(rows)
            if attrs:
                if 'power_domain' not in columns:
                    columns['power_domain'] = len(adjusted[0])
                    adjusted[0].append('POWER_DOMAIN')
                col = columns['power_domain']
                for row in adjusted[1:]:
                    if len(row) <= columns['signal_name']:
                        continue
                    signal = by_name.get(row[columns['signal_name']].casefold())
                    if signal:
                        row.extend([''] * max(0, col + 1 - len(row)))
                        row[col] = signal.power_domain
            details = []
            for a, b in zip(original, effective):
                entry = {rows[0][index]: rows[a.row_order + 1][index] if index < len(rows[a.row_order + 1]) else '' for index in range(len(rows[0]))}
                entry.update(signal=a.name, original_power_domain=a.power_domain, effective_power_domain=b.power_domain, source_workbook=str(path), source_sheet=sheet)
                details.append(entry)
            signals[module], sheets[master_sheet] = details, adjusted
            changed = module in old and old[module].get('sha256') != before
            mappings.append(dict(module=module, source_workbook=str(path), source_sheet=sheet, column=column, master_sheet=master_sheet, sha256=before, changed=changed, status='PASS'))
            if changed:
                issues.append(dict(module=module, sheet=sheet, signal='', status='Warning', message='原始文件已变化，需要更新有效 Master'))
            if file_sha256(path) != before:
                raise ValueError('读取期间原始文件发生变化，请重试')
        except Exception as exc:
            issues.append(dict(module=module, sheet=sheet, signal=signal_context, status='Error', message=f'MSI 检查失败：{exc}'))
    identity = [{k: v for k, v in row.items() if k not in {'changed', 'status'}} for row in mappings]
    fingerprint = digest(dict(sources=identity, attributes=raw['attribute_overrides'], case_overrides=raw['overrides'], sheets=sheets))
    return EffectiveMSI(raw, mappings, signals, sheets, issues, fingerprint)


def workbook_diff(path, report):
    import difflib
    try:
        actual = read_workbook(path)
    except Exception as exc:
        return f'无法读取人工修改版本：{exc}'
    lines = []
    for sheet in sorted(set(report.sheets) | {s for s in actual if not s.startswith('_')}):
        render = lambda rows: [' | '.join(str(x) for x in row) + '\n' for row in rows]
        lines.extend(difflib.unified_diff(render(actual.get(sheet, [])), render(report.sheets.get(sheet, [])), fromfile=f'现有/{sheet}', tofile=f'自动重建/{sheet}'))
    return ''.join(lines) or '工作簿格式或追溯工作表存在修改。'


def prepare(payload, workflow, project_dir, *, restore=False):
    report = inspect(payload, workflow)
    report.require_valid()
    state = workflow.setdefault('msi_master', {})
    if state.get('source_root') and 'legacy_sources' not in state:
        state['legacy_sources'] = copy.deepcopy(state.get('sources', []))
    previous = Path(state.get('path') or project_dir / 'msi' / 'MSI_Master.xlsx')
    manual = previous.is_file() and (not state.get('sha256') or file_sha256(previous) != state['sha256'])
    if manual and not restore:
        raise ProjectValidationError('MSI_MASTER_MANUAL_EDIT', '检测到 Master 人工修改，未覆盖。请在“预览有效 MSI”检查差异并恢复自动版本；人工属性调整请录入第 9 节。\n' + workbook_diff(previous, report))
    if manual:
        backup = project_dir / 'msi' / 'manual_backups' / file_sha256(previous) / 'MSI_Master.xlsx'
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            shutil.copyfile(previous, backup)
    root = project_dir / 'msi' / 'versions' / report.fingerprint
    target = root / 'MSI_Master.xlsx'
    # Never overwrite a changed historical snapshot. Rebuild into a new revision.
    if target.exists() and (manual or (state.get('fingerprint') != report.fingerprint)):
        root = Path(tempfile.mkdtemp(prefix=report.fingerprint[:12] + '_', dir=root.parent))
        target = root / 'MSI_Master.xlsx'
    if not target.exists():
        root.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        book.remove(book.active)
        for name, rows in report.sheets.items():
            sheet = book.create_sheet(name)
            for row in rows:
                sheet.append(row)
        for title, records in (('_Source_Info', report.mappings), ('_Attribute_Overrides', report.payload['attribute_overrides']), ('_Case_Overrides', report.payload['overrides'])):
            sheet = book.create_sheet(title)
            if records:
                headers = list(records[0])
                sheet.append(headers)
                for record in records:
                    sheet.append([str(record.get(k, '')) for k in headers])
        fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=root)
        os.close(fd)
        try:
            book.save(temporary)
            for row in report.mappings:
                if file_sha256(Path(row['source_workbook'])) != row['sha256']:
                    raise ProjectValidationError('MSI_SOURCE_CHANGED', '构建期间 MSI 来源变化，请重新生成')
            os.replace(temporary, target)
        finally:
            unlink_missing_ok(Path(temporary))
            book.close()
    sha = file_sha256(target)
    manifest = dict(fingerprint=report.fingerprint, sha256=sha, sources=report.mappings, attribute_overrides=report.payload['attribute_overrides'], case_overrides=report.payload['overrides'])
    (root / 'source_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    history = state.setdefault('history', [])
    if not history or history[-1].get('sha256') != sha:
        history.append(dict(path=str(target), sha256=sha, fingerprint=report.fingerprint, version=len(history) + 1, created_at=datetime.now(timezone.utc).isoformat()))
    state.update(path=str(target), sha256=sha, fingerprint=report.fingerprint, sources=report.mappings, version=history[-1]['version'])
    effective = copy.deepcopy(report.payload)
    effective['roots']['msi'] = str(root)
    effective['msi_mappings'] = [dict(module=row['module'], file=target.name, sheet=row['master_sheet'], column=row['column']) for row in report.mappings]
    effective['attribute_overrides'] = []  # Already materialized in the Master.
    return effective, target, manifest


def snapshot_for_generation(effective, master, manifest, output):
    root = output / '_inputs' / manifest['sha256']
    root.mkdir(parents=True, exist_ok=True)
    target = root / master.name
    if target.exists() and file_sha256(target) != manifest['sha256']:
        raise ProjectValidationError('MSI_SNAPSHOT_CHANGED', f'历史生成输入被修改，禁止覆盖：{target}')
    if not target.exists():
        shutil.copyfile(master, target)
    if file_sha256(target) != manifest['sha256']:
        raise ProjectValidationError('MSI_SNAPSHOT_CHANGED', '复制期间 Master 发生变化')
    (root / 'source_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    effective = copy.deepcopy(effective)
    effective['roots']['msi'] = str(root)
    return effective, target
