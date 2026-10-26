from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ._compat37 import unlink_missing_ok
from ibis_auto.config import parse_config

from .config_adapter import core_semantics, payload_to_core_config
from .errors import ProjectStorageError, ProjectValidationError


SECTION_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")


def _section(sheet, title: str) -> None:
    sheet.append([title])
    cell = sheet.cell(sheet.max_row, 1)
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = SECTION_FILL


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL


def export_config_excel(payload: dict[str, Any], target: Path) -> Path:
    """Export complete Section 1-10 data without touching the source workbook."""

    target = target.expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IBIS_Config"

    _section(sheet, "1. Project / Path")
    for label, key in (("Model Root", "model"), ("MSI Root", "msi"), ("SPF Root", "spf"), ("T2B Template Root", "template")):
        sheet.append([label, str(payload["roots"].get(key, ""))])

    _section(sheet, "2. Module / Direction / Ron")
    impedances = sorted({int(value) for module in payload["modules"] for value in module["impedances"]})
    _header(sheet, [
        "Module", "Generate Type", *[f"{value}ohm" for value in impedances],
        "IBIS IO Voltage Domain", "IBIS VIH Voltage Domain",
        "Tr MAX", "Tr TYP", "Tr MIN", "Tf MAX", "Tf TYP", "Tf MIN",
    ])
    for module in payload["modules"]:
        selected = {int(value) for value in module["impedances"]}
        sheet.append([
            module["name"], module["generate_type"],
            *["Y" if value in selected else "" for value in impedances],
            module["ibis_io_voltage_domain"], module["ibis_vih_voltage_domain"],
            *[module["tr"][name] for name in ("MAX", "TYP", "MIN")],
            *[module["tf"][name] for name in ("MAX", "TYP", "MIN")],
        ])

    _section(sheet, "3.1 Process / Temperature")
    _header(sheet, ["Parameter", "MAX", "TYP", "MIN"])
    for field, label in (("process", "Process"), ("temperature", "Temperature")):
        sheet.append([label, *[payload["corners"][name][field] for name in ("MAX", "TYP", "MIN")]])

    _section(sheet, "3.2 Voltage Domain")
    _header(sheet, ["Voltage Domain", "MAX", "TYP", "MIN"])
    for domain in payload["voltage_order"]:
        sheet.append([domain, *[payload["corners"][name]["voltages"].get(domain, "") for name in ("MAX", "TYP", "MIN")]])

    _section(sheet, "3.3 Calibration")
    _header(sheet, ["Signal", "MAX", "TYP", "MIN"])
    calibration_order: list[str] = []
    for name in ("MAX", "TYP", "MIN"):
        for signal in payload["corners"][name]["calibration"]:
            if signal not in calibration_order:
                calibration_order.append(signal)
    for signal in calibration_order:
        sheet.append([signal, *[payload["corners"][name]["calibration"].get(signal, "") for name in ("MAX", "TYP", "MIN")]])

    _section(sheet, "3.4 SPF Match")
    _header(sheet, ["Parameter", "MAX", "TYP", "MIN"])
    sheet.append(["SPF Match Key", *[payload["corners"][name]["spf_key"] for name in ("MAX", "TYP", "MIN")]])

    _section(sheet, "4. Simulation Option")
    _header(sheet, ["Simulation Option"])
    for option in payload["simulation_options"]:
        sheet.append([option])

    _section(sheet, "5. IBIS Pin Mapping")
    _header(sheet, ["Module", "Direction", "Original Pin", "", "", "", "IBIS Alias"])
    for row in payload["aliases"]:
        sheet.append([row["module"], row["direction"], row["original_pin"], "", "", "", row["alias"]])

    _section(sheet, "6. CKT Append / Reference")
    _header(sheet, ["Module", "Direction", "SPICE Body"])
    for row in payload["appends"]:
        sheet.append([row["module"], row["direction"], row["body"]])

    _section(sheet, "7. T2B Template Mapping")
    _header(sheet, ["Module", "Direction", "Template Type"])
    for (module, direction), template_type in payload_to_core_config(payload).templates.items():
        sheet.append([module, direction, template_type])

    _section(sheet, "8. Signal Override")
    _header(sheet, ["Module", "Direction", "Corner", "Signal", "Value"])
    for row in payload["overrides"]:
        sheet.append([row.get(key, "") for key in ("module", "direction", "corner", "signal", "value")])

    _section(sheet, "9. MSI Attribute Override")
    _header(sheet, ["Module", "Signal", "Attribute", "Value"])
    for row in payload["attribute_overrides"]:
        sheet.append([row.get(key, "") for key in ("module", "signal", "attribute", "value")])

    _section(sheet, "10. MSI Mapping")
    _header(sheet, ["Module", "Workbook", "Sheet", "PIN_NAME Column"])
    for row in payload["msi_mappings"]:
        sheet.append([row.get(key, "") for key in ("module", "file", "sheet", "column")])

    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            80, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )

    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix=target.name + ".", suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        os.replace(temporary, target)
    except OSError as exc:
        if temporary:
            unlink_missing_ok(temporary)
        raise ProjectStorageError("CONFIG_EXPORT_FAILED", f"Config Excel 导出失败：{target}（{exc}）") from exc
    return target


def verify_export_round_trip(payload: dict[str, Any], path: Path) -> None:
    expected = core_semantics(payload_to_core_config(payload))
    actual = core_semantics(parse_config(path))
    if actual != expected:
        raise ProjectValidationError(
            "CONFIG_EXPORT_SEMANTIC_MISMATCH",
            "导出的 Config Excel 重新解析后与当前 ProjectStore 语义不一致",
        )
