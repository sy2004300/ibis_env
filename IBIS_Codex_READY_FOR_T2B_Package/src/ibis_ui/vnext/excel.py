from __future__ import annotations
import copy
import json
import os
import re
import tempfile
from pathlib import Path, PureWindowsPath
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side, PatternFill, Font
from .model import empty, validate_shape, CORNERS, SOURCES
from ..errors import ExcelImportError

TABLES = {
    'modules': ('02 Resolved Module Plans', [('name', 'Module'), ('generate_type', 'Generate Type'), ('topology', 'Topology'), ('io_domain', 'IO Domain'), ('vih_domain', 'VIH Domain'), ('ron', 'Ron (Ω)')]),
    'pins': ('04 IBIS Pin Mapping', [('module', 'Module'), ('direction', 'Direction'), ('t2b_pin', 'T2B Pin'), ('spf_pin', 'SPF Pin'), ('status', 'Status')]),
    'appends': ('05 CKT Append', [('module', 'Module'), ('direction', 'Direction'), ('body', 'SPICE Body')]),
    'overrides': ('06 Signal Override', [('module', 'Module'), ('direction', 'Direction'), ('signal', 'Signal'), ('MAX', 'MAX'), ('TYP', 'TYP'), ('MIN', 'MIN')]),
    'msi_mappings': ('07 MSI Mapping', [('module', 'Module'), ('file', 'Workbook'), ('sheet', 'Sheet'), ('column', 'PIN_NAME Column'), ('status', 'Status')]),
}


def text(value):
    return '' if value is None else str(value)


def sections(sheet):
    current = None
    result = {}
    for raw in sheet.iter_rows(values_only=True):
        row = [text(v) for v in raw]
        first = row[0].strip() if row else ''
        if re.match(r'^0[1-7]\s+', first):
            current = first[2:].strip().split('/')[0].strip().casefold()
            result[current] = []
        elif current and any(row):
            if not any(row[1:]) and (len(first) > 45 or first.startswith(('只预留', '可以预先', '仅填写', '下一步', '示例：', '新版配置'))):
                continue
            result[current].append(row)
    return result


def table(rows):
    if not rows:
        return []
    headers = [h.strip() for h in rows[0]]
    result = []
    for row in rows[1:]:
        if not any(v.strip() for v in row):
            continue
        result.append({h: row[i] if i < len(row) else '' for i, h in enumerate(headers) if h})
    return result


def detect(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Input' in book.sheetnames and 'input sources' in sections(book['Input']):
            return 'input'
        if 'Config_Generated' in book.sheetnames and 'resolved sources' in sections(book['Config_Generated']):
            return 'generated'
        if 'Generated Demo' in book.sheetnames:
            return 'demo'
        return 'legacy'
    finally:
        book.close()


def read(path):
    kind = detect(path)
    if kind == 'legacy':
        return None
    data = empty()
    if kind == 'demo':
        data['provenance']['notice'] = 'Generated Demo 仅为布局示例，未读取为生产配置。请导入 Input 或正式 Config_Generated。'
        return data
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        blocks = sections(book['Input' if kind == 'input' else 'Config_Generated'])
        metadata = {}
        if '_VNext' in book.sheetnames:
            parts = [row[1] for row in book["_VNext"].iter_rows(values_only=True) if row and row[0] in ("metadata", "metadata_part")]
            metadata = json.loads("".join(parts)) if parts else {}
            if metadata.get('schema_version', 1) != 1:
                raise ExcelImportError('VNEXT_EXCEL_VERSION', '完整 Excel 版本不受支持')
        sources = table(blocks.get('input sources' if kind == 'input' else 'resolved sources', []))
        for row in sources:
            if row.get('Source') in SOURCES:
                value = row.get('Path / File', row.get('Resolved Path / File', '')).strip()
                if value and not Path(value).is_absolute() and not PureWindowsPath(value).is_absolute():
                    value = str((Path(path).parent / value).resolve())
                data['sources'][SOURCES[row['Source']]] = value
        modules = table(blocks.get('module configuration' if kind == 'input' else 'resolved module plans', []))
        for row in modules:
            name = row.get('Module', row.get('Module (SPF 匹配名)', '')).strip()
            if name:
                data['modules'].append(dict(name=name, generate_type=row.get('Generate Type', row.get('Direction', '')).strip().upper(), topology=row.get('Topology', '').strip().upper(), io_domain=row.get('IO Domain', '').strip(), vih_domain=row.get('VIH Domain', '').strip(), ron=row.get('Ron (Ω)', '').strip()))
        fields = {'Process': 'process', 'Temperature (°C)': 'temperature', 'Temperature': 'temperature', 'Tr': 'tr', 'Tf': 'tf', 'SPF Match Key': 'spf_key', 'Match Key': 'spf_key'}
        if kind == 'input':
            for row in table(blocks.get('corner configuration', [])):
                key = fields.get(row.get('Parameter'))
                if key:
                    for c in CORNERS:
                        data['corners'][c][key] = row.get(c, '').strip()
            for section, target, field, header in [('optional voltage values', 'voltages', 'domain', 'Domain'), ('optional global signal defaults', 'signals', 'signal', 'Signal')]:
                for row in table(blocks.get(section, [])):
                    if row.get(header, '').strip():
                        data[target].append(dict({field: row[header].strip(), 'source': 'INPUT'}, **{c: row.get(c, '').strip() for c in CORNERS}))
        else:
            for row in table(blocks.get('corner', [])):
                category, name = row.get('Category', ''), row.get('Parameter', '')
                if category in ('Corner', 'Timing', 'SPF') and name in fields:
                    for c in CORNERS:
                        data['corners'][c][fields[name]] = row.get(c, '').strip()
                elif category in ('Voltage', 'Signal') and name:
                    field = 'domain' if category == 'Voltage' else 'signal'
                    target = 'voltages' if category == 'Voltage' else 'signals'
                    data[target].append(dict({field: name, 'source': row.get('Source', 'MANUAL')}, **{c: row.get(c, '').strip() for c in CORNERS}))
            for field, (title, columns) in TABLES.items():
                if field == 'modules':
                    continue
                key = title[3:].casefold()
                for row in table(blocks.get(key, [])):
                    record = {name: row.get(label, '') for name, label in columns}
                    if record.get('module'):
                        if field in ('pins', 'msi_mappings') and record.get('status', '').startswith('='):
                            record['status'] = 'UNRESOLVED'
                        data[field].append(record)
            data['provenance'] = metadata.get('provenance', {})
            # Visible cells are authoritative; metadata only retains fingerprints.
            for field in ('pins', 'msi_mappings'):
                old = metadata.get(field, [])
                keys = ('module', 'direction', 't2b_pin') if field == 'pins' else ('module',)
                for row in data[field]:
                    previous = next((r for r in old if all(r.get(k) == row.get(k) for k in keys)), {})
                    changed = any(row.get(k, '') != previous.get(k, '') for k in row if k != 'status')
                    confirmed = row.get('status') == 'MANUAL' and previous.get('status') == 'STALE'
                    row['fingerprint'] = '' if changed or confirmed else previous.get('fingerprint', '')
                    if changed and any(row.get(k) for k in (('spf_pin',) if field == 'pins' else ('file', 'sheet', 'column'))):
                        row['status'] = 'MANUAL'
        data['provenance']['import_format'] = kind
        validate_shape(data)
        return data
    finally:
        book.close()


def write(data, target):
    validate_shape(data)  # Storage shape only: missing values MUST remain exportable.
    target = Path(target).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Config_Generated'
    title_fill = PatternFill('solid', fgColor='204A70')
    header_fill = PatternFill('solid', fgColor='DCE8F2')
    border = Border(**{k: Side(style='thin', color='A0A0A0') for k in ('left', 'right', 'top', 'bottom')})
    def append(values, style=None):
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.border = border
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if style:
                cell.fill = style
                cell.font = Font(bold=True, color='FFFFFF' if style == title_fill else '000000')
    def section(title, headers):
        append([title], title_fill)
        append(headers, header_fill)
    section('01 Resolved Sources', ['Source', 'Resolved Path / File', 'Status'])
    for label, key in SOURCES.items():
        append([label, data['sources'].get(key, ''), 'MANUAL' if data['sources'].get(key) else ('AUTO_MATCH' if key == 'template' else 'UNRESOLVED')])
    title, columns = TABLES['modules']
    section(title, [label for _, label in columns])
    for row in data['modules']:
        append([row.get(key, '') for key, _ in columns])
    section('03 Corner / Power / Signal', ['Category', 'Parameter', 'MAX', 'TYP', 'MIN', 'Source'])
    for category, label, key in [('Corner', 'Process', 'process'), ('Corner', 'Temperature', 'temperature'), ('Timing', 'Tr', 'tr'), ('Timing', 'Tf', 'tf'), ('SPF', 'Match Key', 'spf_key')]:
        append([category, label] + [data['corners'][c].get(key, '') for c in CORNERS] + ['MANUAL'])
    for field, category, name in [('voltages', 'Voltage', 'domain'), ('signals', 'Signal', 'signal')]:
        for row in data[field]:
            append([category, row.get(name, '')] + [row.get(c, '') for c in CORNERS] + [row.get('source', 'MANUAL')])
    for field, (title, columns) in TABLES.items():
        if field == 'modules':
            continue
        section(title, [label for _, label in columns])
        for row in data[field]:
            append([row.get(key, '') for key, _ in columns])
    report = book.create_sheet('检查结果')
    report.append(['级别', '类别', '模块', '方向', '工艺角', '字段', '文件/工作表', '原因'])
    for row in data.get('issues', []):
        report.append([row.get(k, '') for k in ('severity', 'code', 'module', 'direction', 'corner', 'field', 'location', 'message')])
    for row in report:
        for cell in row:
            cell.border = border
            if cell.row == 1:
                cell.fill = header_fill
    metadata = book.create_sheet('_VNext')
    metadata.sheet_state = 'hidden'
    # Split before assigning cells: openpyxl truncates individual values to
    # Excel's 32767-character limit at assignment time.
    encoded = json.dumps({k: copy.deepcopy(data[k]) for k in ('schema_version', 'provenance', 'pins', 'msi_mappings')}, ensure_ascii=False)
    metadata.append(['metadata', encoded[:30000]])
    for index in range(30000, len(encoded), 30000):
        metadata.append(['metadata_part', encoded[index:index + 30000]])
    for col in ('A','B','C','D','E','F'):
        sheet.column_dimensions[col].width = 28 if col != 'B' else 48
    sheet.freeze_panes = 'A3'
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=str(target.parent), suffix='.xlsx', delete=False) as stream:
            temporary = Path(stream.name)
        book.save(str(temporary))
        os.replace(str(temporary), str(target))
    finally:
        book.close()
        if temporary and temporary.exists():
            temporary.unlink()
    return target
