from __future__ import annotations

import hashlib
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from ibis_auto.errors import IbisError
from ibis_auto.msi import _columns
from ibis_auto.ooxml import read_workbook

from .errors import ProjectStorageError, ProjectValidationError


MASTER_FILENAME = "MSI_Master.xlsx"
EDITABLE_COLUMNS = {"POWERDOMAIN", "DEFAULTVALUE", "DEFAULTVALUEPHYRESET", "DEFAULTVALUEPOWEROFF"}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normal(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).casefold())


def _safe_sheet_name(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", value).strip(" '") or "MSI"
    base = base[:31]
    name = base
    counter = 2
    while name.casefold() in used or name.casefold() in {"_msi_mapping", "_source_info"}:
        suffix = f"_{counter}"
        name = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(name.casefold())
    return name


@dataclass(frozen=True)
class MSISource:
    path: Path
    sheet: str
    pin_column: str
    sha256: str
    rows: int


@dataclass(frozen=True)
class MergePreview:
    sources: tuple[MSISource, ...]
    excluded: tuple[str, ...]
    mappings: tuple[dict[str, str], ...]
    ambiguous_modules: tuple[str, ...]
    differences: tuple[str, ...]


def scan_msi_sources(root: Path) -> tuple[list[MSISource], list[str]]:
    if not root.is_dir():
        raise ProjectValidationError("MSI_ROOT_MISSING", f"MSI Root 不存在：{root}")
    sources: list[MSISource] = []
    excluded: list[str] = []
    for path in sorted(root.rglob("*.xlsx")):
        if path.name.casefold() == MASTER_FILENAME.casefold() or path.name.startswith("~$"):
            excluded.append(f"{path}：Master/临时文件不作为来源")
            continue
        sha256 = file_sha256(path)
        try:
            sheets = read_workbook(path)
        except Exception as exc:
            excluded.append(f"{path}：无法读取（{exc}）")
            continue
        for sheet, rows in sheets.items():
            if not rows:
                excluded.append(f"{path.name}/{sheet}：空 Sheet")
                continue
            try:
                columns = _columns(rows[0])
            except IbisError as exc:
                excluded.append(f"{path.name}/{sheet}：非有效 MSI（{exc}）")
                continue
            sources.append(MSISource(path.resolve(), sheet, rows[0][columns["signal_name"]], sha256, max(0, len(rows) - 1)))
    return sources, excluded


def _automatic_mappings(modules: list[str], sources: list[MSISource]) -> tuple[list[dict[str, str]], list[str]]:
    mappings: list[dict[str, str]] = []
    ambiguous: list[str] = []
    for module in modules:
        key = _normal(module)
        matches = [source for source in sources if key in {_normal(source.path.stem), _normal(source.sheet)}]
        if len(matches) == 1:
            source = matches[0]
            mappings.append({
                "module": module,
                "source_file": str(source.path),
                "source_sheet": source.sheet,
                "column": source.pin_column,
            })
        else:
            ambiguous.append(module)
    return mappings, ambiguous


def _existing_edit_diff(master: Path, sources: list[MSISource]) -> list[str]:
    if not master.is_file():
        return []
    try:
        workbook = load_workbook(master, read_only=True, data_only=False)
    except Exception as exc:
        return [f"现有 Master 无法读取：{exc}"]
    differences: list[str] = []
    source_by_sheet = {source.sheet.casefold(): source for source in sources}
    for sheet in workbook.worksheets:
        if sheet.title.startswith("_"):
            continue
        source = source_by_sheet.get(sheet.title.casefold())
        if source is None:
            differences.append(f"现有 Master Sheet {sheet.title} 不再有直接同名来源")
            continue
        source_rows = read_workbook(source.path).get(source.sheet, [])
        master_rows = [[str(cell.value or "").strip() for cell in row] for row in sheet.iter_rows()]
        if master_rows != source_rows:
            differences.append(f"{sheet.title}：Master 内容与来源存在修改")
    workbook.close()
    return differences


def preview_merge(root: Path, modules: list[str], master: Path | None = None) -> MergePreview:
    sources, excluded = scan_msi_sources(root)
    mappings, ambiguous = _automatic_mappings(modules, sources)
    differences = _existing_edit_diff(master, sources) if master else []
    return MergePreview(tuple(sources), tuple(excluded), tuple(mappings), tuple(ambiguous), tuple(differences))


def _copy_sheet(source: MSISource, target, preserved: dict[tuple[str, str], Any] | None = None) -> None:
    workbook = load_workbook(source.path, read_only=True, data_only=False)
    sheet = workbook[source.sheet]
    rows = list(sheet.iter_rows(values_only=True))
    for row in rows:
        target.append(list(row))
    if preserved and rows:
        headers = [str(value or "") for value in rows[0]]
        normalized = [_normal(value) for value in headers]
        pin_indexes = [i for i, name in enumerate(normalized) if name in {"pinname", "name"}]
        if len(pin_indexes) == 1:
            pin_index = pin_indexes[0]
            editable = [i for i, name in enumerate(normalized) if name.upper() in EDITABLE_COLUMNS]
            for row_index in range(2, target.max_row + 1):
                signal = str(target.cell(row_index, pin_index + 1).value or "")
                for column_index in editable:
                    key = (signal.casefold(), normalized[column_index])
                    if key in preserved:
                        target.cell(row_index, column_index + 1).value = preserved[key]
    workbook.close()


def _preserved_values(master: Path) -> dict[str, dict[tuple[str, str], Any]]:
    if not master.is_file():
        return {}
    workbook = load_workbook(master, read_only=True, data_only=False)
    result: dict[str, dict[tuple[str, str], Any]] = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if sheet.title.startswith("_") or not rows:
            continue
        headers = [_normal(value or "") for value in rows[0]]
        pin_indexes = [i for i, name in enumerate(headers) if name in {"pinname", "name"}]
        if len(pin_indexes) != 1:
            continue
        editable = [i for i, name in enumerate(headers) if name.upper() in EDITABLE_COLUMNS]
        values: dict[tuple[str, str], Any] = {}
        for row in rows[1:]:
            signal = str(row[pin_indexes[0]] or "") if pin_indexes[0] < len(row) else ""
            for index in editable:
                if index < len(row):
                    values[(signal.casefold(), headers[index])] = row[index]
        result[sheet.title.casefold()] = values
    workbook.close()
    return result


def merge_msi_master(
    preview: MergePreview,
    target: Path,
    *,
    explicit_mappings: list[dict[str, str]] | None = None,
    preserve_edits: bool = True,
) -> tuple[Path, list[dict[str, str]], list[dict[str, str]]]:
    if not preview.sources:
        raise ProjectValidationError("MSI_SOURCE_EMPTY", "没有可合并的有效 MSI Sheet")
    mappings = list(explicit_mappings if explicit_mappings is not None else preview.mappings)
    if preview.ambiguous_modules and explicit_mappings is None:
        raise ProjectValidationError(
            "MSI_MAPPING_AMBIGUOUS",
            f"以下 Module 的 MSI 来源不唯一，必须显式选择：{', '.join(preview.ambiguous_modules)}",
        )
    target = target.expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    source_hashes = {source.path: file_sha256(source.path) for source in preview.sources}
    preserved = _preserved_values(target) if preserve_edits else {}
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    source_to_master: dict[tuple[str, str], str] = {}
    source_info: list[dict[str, str]] = []
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for index, source in enumerate(preview.sources, 1):
        proposed = source.sheet if sum(item.sheet.casefold() == source.sheet.casefold() for item in preview.sources) == 1 else f"{source.path.stem}_{source.sheet}"
        master_sheet = _safe_sheet_name(proposed, used)
        target_sheet = workbook.create_sheet(master_sheet)
        _copy_sheet(source, target_sheet, preserved.get(master_sheet.casefold()))
        source_to_master[(str(source.path), source.sheet)] = master_sheet
        source_info.append({
            "source_workbook": str(source.path), "source_sheet": source.sheet,
            "master_sheet": master_sheet, "sha256": source.sha256,
            "imported_at": now, "version": str(index),
        })
    mapping_sheet = workbook.create_sheet("_MSI_Mapping")
    mapping_sheet.append(["Module", "Master Workbook", "Master Sheet", "PIN_NAME Column"])
    master_mappings: list[dict[str, str]] = []
    for row in mappings:
        key = (str(Path(row["source_file"]).resolve()), str(row["source_sheet"]))
        if key not in source_to_master:
            raise ProjectValidationError("MSI_MAPPING_SOURCE_INVALID", f"Module {row['module']} 的 MSI 来源不在本次合并预览中")
        mapped = {
            "module": str(row["module"]), "file": MASTER_FILENAME,
            "sheet": source_to_master[key], "column": str(row["column"]),
        }
        master_mappings.append(mapped)
        mapping_sheet.append([mapped["module"], mapped["file"], mapped["sheet"], mapped["column"]])
    info_sheet = workbook.create_sheet("_Source_Info")
    info_headers = ["Source Workbook", "Source Sheet", "Master Sheet", "SHA256", "Imported At", "Version"]
    info_sheet.append(info_headers)
    for row in source_info:
        info_sheet.append([row[key] for key in ("source_workbook", "source_sheet", "master_sheet", "sha256", "imported_at", "version")])

    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix=target.stem + ".", suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        load_workbook(temporary, read_only=True).close()
        for source, before in source_hashes.items():
            if file_sha256(source) != before:
                raise ProjectStorageError("MSI_SOURCE_CHANGED", f"合并期间原始 MSI 发生变化：{source}")
        os.replace(temporary, target)
    except Exception as exc:
        if temporary:
            temporary.unlink(missing_ok=True)
        if isinstance(exc, (ProjectStorageError, ProjectValidationError)):
            raise
        raise ProjectStorageError("MSI_MASTER_SAVE_FAILED", f"MSI Master 保存失败，旧文件保持不变：{target}（{exc}）") from exc
    return target, master_mappings, source_info
