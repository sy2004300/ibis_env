"""Synthetic, temporary VNext inputs. Frozen inputs are only read/copied."""
from pathlib import Path
import shutil
from openpyxl import load_workbook
from ibis_ui.vnext.scan import precheck
from ibis_ui.vnext.model import empty, CORNERS

PACKAGE = Path(__file__).resolve().parents[1]
TEMPLATES = PACKAGE/'templates/t2b'
MODULE = 'demo_phy_io_diff'


def fixture(root, se=False):
    source = root/'sources'
    shutil.copytree(str(PACKAGE/'inputs/msi'),str(source/'msi'))
    shutil.copytree(str(PACKAGE/'inputs/spf'),str(source/'spf'))
    shutil.copytree(str(TEMPLATES),str(source/'templates'))
    models = source/'models'
    models.mkdir()
    for name, suffixes in [('rbd',('f','t','s','fs','sf')), ('cap',('f','t','s','fs','sf')), ('mos_var',('ff','tt','ss','fs','sf'))]:
        (models/(name+'.lib')).write_text('* Synthetic sections for VNext tests, not a real PDK\n'+''.join('.lib %s_%s\n.endl %s_%s\n' % (name,s,name,s) for s in suffixes),encoding='utf-8')
    data = empty()
    data['sources'] = dict(model=str(models),msi=str(source/'msi'),spf=str(source/'spf'),template=str(source/'templates'))
    data['modules'] = [dict(name=MODULE,generate_type='BOTH',topology='DIFF',io_domain='VDDQ',vih_domain='VDD',ron='30,34,40')]
    for c,p,t,k in zip(CORNERS,('FF','TT','SS'),('125','25','-40'),('rcbest','typical','rcworst')):
        data['corners'][c] = dict(process=p,temperature=t,tr='20p',tf='20p',spf_key=k)
    data['voltages'] = [dict(domain=d,**dict(zip(CORNERS,values))) for d,values in [('VDD',('0.88','0.8','0.72')),('VDDQ',('0.88','0.8','0.72')),('VSS',('0','0','0'))]]
    if se:
        name='demo_phy_io_se'
        module=dict(data['modules'][0],name=name,topology='SE',generate_type='DRV',ron='60,80')
        data['modules'].append(module)
        for path in (source/'spf').glob('*.spf'):
            text=path.read_text().replace(MODULE,name).replace('IOPADT','IOPAD').replace(' IOPADC','')
            (path.parent/path.name.replace(MODULE,name)).write_text(text)
        path=source/'msi/demo_msi.xlsx'
        book=load_workbook(path)
        sheet=book.copy_worksheet(book[MODULE]);sheet.title=name
        for row in list(sheet):
            if row[0].value=='IOPADT': row[0].value='IOPAD'
            if row[0].value=='IOPADC':
                for cell in row: cell.value=None
        book.save(path);book.close()
    return data


def confirmed(data):
    data=precheck(data,TEMPLATES)
    for row in data['pins']:
        if not row['spf_pin']:
            row['spf_pin']={'ipadt':'IOPADT','ipadc':'IOPADC','iopad':'IOPAD','ipad':'IOPAD'}.get(row['t2b_pin'],'')
            row['status']='MANUAL'
    return precheck(data,TEMPLATES)


def input_book(data,path):
    book=load_workbook(PACKAGE/'docs/vnext/IBIS_Config_Input_Final.xlsx')
    sheet=book['Input']
    # Populate by header/title, not fixed row offsets; extend rows as engineers do.
    for row in sheet:
        if row[0].value in ('Model Source','MSI Source','SPF Source','T2B Source'):
            key={'Model Source':'model','MSI Source':'msi','SPF Source':'spf','T2B Source':'template'}[row[0].value]
            sheet.cell(row[0].row,2,data['sources'][key])
    for title,rows in [('02',[[m[k] for k in ('name','generate_type','topology','io_domain','vih_domain','ron')] for m in data['modules']]),('04',[[r['domain']]+[r[c] for c in CORNERS] for r in data['voltages']]),('05',[[r['signal']]+[r.get(c,'') for c in CORNERS] for r in data['signals']])]:
        start=next(row[0].row for row in sheet if str(row[0].value).startswith(title+' '))+2
        sheet.insert_rows(start,len(rows))
        for i,values in enumerate(rows):
            for j,v in enumerate(values,1): sheet.cell(start+i,j,v)
    labels={'Process':'process','Temperature (°C)':'temperature','Tr':'tr','Tf':'tf','SPF Match Key':'spf_key'}
    for row in sheet:
        if row[0].value in labels:
            for j,c in enumerate(CORNERS,2): sheet.cell(row[0].row,j,data['corners'][c][labels[row[0].value]])
    book.save(path);book.close()
    return path
