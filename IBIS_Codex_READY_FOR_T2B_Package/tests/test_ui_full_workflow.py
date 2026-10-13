from __future__ import annotations

import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ibis_auto.config import parse_config
from ibis_ui.config_adapter import core_semantics, payload_to_core_config
from ibis_ui.controller import ProjectController
from ibis_ui.excel_importer import ExcelImporter
from ibis_ui.models import PROJECT_SCHEMA_VERSION, ProjectDocument
from ibis_ui.msi_master import MASTER_FILENAME, merge_msi_master, preview_merge
from ibis_ui.project_store import ProjectStore
from ibis_ui.runtime import ResourceResolver
from ibis_ui.working_templates import validate_template_text


ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "inputs" / "config" / "ibis_config.xlsx"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class FullWorkflowTests(unittest.TestCase):
    def test_schema_1_migrates_without_losing_config(self):
        original = ExcelImporter().import_project(CONFIG, "Legacy")
        payload = original.to_dict()
        payload["schema_version"] = 1
        payload.pop("workflow")
        migrated = ProjectDocument.from_dict(payload)
        self.assertEqual(migrated.schema_version, PROJECT_SCHEMA_VERSION)
        self.assertEqual(migrated.config, original.config)
        self.assertIn("working_templates", migrated.workflow)
        self.assertIn("generation", migrated.workflow)

    def test_config_section_1_to_10_export_round_trip_and_source_immutable(self):
        before = sha256(CONFIG)
        with tempfile.TemporaryDirectory() as tmp:
            store = ProjectStore(Path(tmp))
            controller = ProjectController(store)
            controller.import_excel(CONFIG)
            controller.replace_config_rows("simulation_options", [".option raw_spacing = keep", ".option runlvl=5"])
            controller.replace_config_rows("appends", [{"module": "demo_phy_io_diff", "direction": "DRV", "body": "Rkeep a b  1k\n+ exact continuation"}])
            target = Path(tmp) / "export.xlsx"
            controller.export_config(target)
            reparsed = parse_config(target)
            expected = payload_to_core_config(controller.active_config)
            self.assertEqual(core_semantics(reparsed), core_semantics(expected))
            self.assertEqual(reparsed.simulation_options[0], ".option raw_spacing = keep")
            self.assertEqual(reparsed.appends[0][2], "Rkeep a b  1k\n+ exact continuation")
        self.assertEqual(sha256(CONFIG), before)

    def test_module_rename_duplicate_delete_updates_references(self):
        with tempfile.TemporaryDirectory() as tmp:
            controller = ProjectController(ProjectStore(Path(tmp)))
            controller.import_excel(CONFIG)
            controller.rename_module("demo_phy_io_diff", "renamed_io")
            for field in ("aliases", "templates", "msi_mappings"):
                self.assertTrue(all(row["module"] == "renamed_io" for row in controller.active_config[field]))
            controller.duplicate_module("renamed_io", "copy_io")
            self.assertIn("copy_io", [row["name"] for row in controller.active_config["modules"]])
            references = controller.delete_module("copy_io", cascade=True)
            self.assertTrue(references)
            self.assertNotIn("copy_io", [row["name"] for row in controller.active_config["modules"]])

    def test_msi_master_multi_workbook_collision_hash_and_preserve_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "sources"
            root.mkdir()
            originals = []
            for filename, power in (("first.xlsx", "VDDQ"), ("second.xlsx", "VDD")):
                path = root / filename
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "same_sheet"
                sheet.append(["PIN_NAME", "WIDTH", "DIRECTION_TYPE", "POWER_DOMAIN", "DEFAULT_VALUE"])
                sheet.append([filename.removesuffix(".xlsx"), 1, "Input", power, 0])
                workbook.save(path)
                originals.append((path, sha256(path)))
            preview = preview_merge(root, [])
            self.assertEqual(len(preview.sources), 2)
            target = Path(tmp) / MASTER_FILENAME
            target, mappings, info = merge_msi_master(preview, target, explicit_mappings=[])
            workbook = load_workbook(target)
            data_sheets = [name for name in workbook.sheetnames if not name.startswith("_")]
            self.assertEqual(len(data_sheets), 2)
            self.assertEqual(len(set(data_sheets)), 2)
            self.assertTrue(all(len(name) <= 31 for name in data_sheets))
            workbook[data_sheets[0]]["D2"] = "EDITED_DOMAIN"
            workbook.save(target)
            second_preview = preview_merge(root, [], target)
            self.assertTrue(second_preview.differences)
            merge_msi_master(second_preview, target, explicit_mappings=[], preserve_edits=True)
            reopened = load_workbook(target, data_only=False)
            self.assertEqual(reopened[data_sheets[0]]["D2"].value, "EDITED_DOMAIN")
            for path, digest in originals:
                self.assertEqual(sha256(path), digest)

    def test_working_template_validation_version_library_and_restore(self):
        with tempfile.TemporaryDirectory() as tmp:
            controller = ProjectController(ProjectStore(Path(tmp)), resources=ResourceResolver.current())
            controller.import_excel(CONFIG)
            first = controller.create_working_template("demo_phy_io_diff", "DRV")
            text = controller.read_working_template("demo_phy_io_diff", "DRV")
            self.assertTrue(validate_template_text(text, "DRV", "DIFF").valid)
            second = controller.save_working_template("demo_phy_io_diff", "DRV", text.replace("[Source] Akrostar", "[Source] Project Edited"))
            self.assertEqual((first["version"], second["version"]), (1, 2))
            service = controller.working_template_service()
            library = service.save_as_library("demo_phy_io_diff", "DRV", "My DRV")
            self.assertTrue(Path(library["path"]).is_file())
            restored = controller.restore_working_template("demo_phy_io_diff", "DRV", 1)
            self.assertEqual(restored["version"], 3)
            self.assertNotIn("Project Edited", controller.read_working_template("demo_phy_io_diff", "DRV"))

    def test_end_to_end_selected_generation_manifest_stale_report_and_reload(self):
        source_hash = sha256(CONFIG)
        with tempfile.TemporaryDirectory() as tmp:
            store = ProjectStore(Path(tmp) / "workspace")
            store.ensure_ready()
            controller = ProjectController(store, resources=ResourceResolver.current())
            controller.import_excel(CONFIG, "Full Workflow")
            controller.update_module("demo_phy_io_diff", "impedances", [30, 40])
            controller.replace_config_rows("simulation_options", [".option method=gear", ".option runlvl=5", ".option ui_e2e=1"])
            aliases = copy.deepcopy(controller.active_config["aliases"])
            next(row for row in aliases if row["original_pin"] == "TXDAT")["alias"] = "txdat_ui"
            controller.replace_config_rows("aliases", aliases)
            controller.create_working_template("demo_phy_io_diff", "DRV")
            text = controller.read_working_template("demo_phy_io_diff", "DRV")
            controller.save_working_template("demo_phy_io_diff", "DRV", text.replace("[Source] Akrostar", "[Source] UI E2E"))
            controller.save_project_draft()

            result = controller.generate_cases([("demo_phy_io_diff", "DRV", 30)])
            self.assertEqual(len(result.cases), 1)
            case = Path(result.cases[0].path)
            self.assertTrue(case.is_dir())
            self.assertEqual(len([path for path in case.iterdir() if path.is_file()]), 7)
            self.assertFalse(case.parent.joinpath("40ohm_ibis").exists())
            self.assertIn(".option ui_e2e=1", (case / "corner_drv_ron30_typ.sp").read_text())
            self.assertIn("txdat_ui", (case / "ckt_topology.sp").read_text())
            self.assertIn("[Source] UI E2E", (case / "demo_phy_io_diff_drv_ron30.t2b").read_text())
            status = controller.generation_statuses()[0]
            self.assertEqual(status["status"], "FAIL")  # expected: modified reference differs from frozen Golden
            self.assertTrue(Path(status["manifest_path"]).is_file())
            line_diff = controller.golden_diff(("demo_phy_io_diff", "DRV", 30))
            self.assertIn("--- Golden/", line_diff)
            self.assertIn("+++ Generated/", line_diff)

            controller.update_corner_voltage("VDDQ", "TYP", "0.51")
            controller.save_project_draft()
            # A Golden mismatch remains FAIL; generate an unvalidated 40-ohm case to exercise STALE.
            controller.generate_cases([("demo_phy_io_diff", "DRV", 40)])
            controller.update_corner_voltage("VDDQ", "TYP", "0.52")
            controller.save_project_draft()
            statuses = {row["ron"]: row["status"] for row in controller.generation_statuses()}
            self.assertEqual(statuses[40], "STALE")
            report_dir = Path(tmp) / "reports"
            csv_path, text_path = controller.export_generation_summary(report_dir)
            self.assertIn("STALE", csv_path.read_text(encoding="utf-8-sig"))
            self.assertIn("Total Cases", text_path.read_text(encoding="utf-8"))

            project_id = controller.project.project_id
            reloaded = ProjectController(store, resources=ResourceResolver.current())
            reloaded.load_project(project_id)
            reloaded_statuses = {row["ron"]: row["status"] for row in reloaded.generation_statuses()}
            self.assertEqual(reloaded_statuses[40], "STALE")
        self.assertEqual(sha256(CONFIG), source_hash)


if __name__ == "__main__":
    unittest.main()
