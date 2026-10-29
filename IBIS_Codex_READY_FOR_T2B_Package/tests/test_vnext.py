from __future__ import annotations
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from ibis_ui.controller import ProjectController
from ibis_ui.project_store import ProjectStore
from ibis_ui.runtime import ResourceResolver
from ibis_ui.models import ProjectDocument
from ibis_ui.errors import ProjectValidationError
from ibis_ui.vnext.model import payload, parse_rons, CORNERS
from ibis_ui.vnext.excel import read, write, detect, sections
from ibis_ui.vnext.scan import precheck, template_pins, hash_file
from ibis_ui.vnext.adapter import core_config
from ibis_ui.vnext.service import generate
from ibis_auto.values import parse_logic
from ibis_auto.spf import parse as parse_spf
from vnext_fixture import fixture, confirmed, input_book, TEMPLATES, MODULE, PACKAGE


class VNextTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.data = fixture(self.root,se=True)
        self.controller = ProjectController(ProjectStore(self.root/'workspace'))

    def tearDown(self):
        self.temp.cleanup()

    def checked(self,data=None):
        return precheck(data or self.data,TEMPLATES)

    def codes(self,data):
        return {r['code'] for r in data['issues'] if r['severity']=='BLOCKING'}

    def assertReady(self,data):
        self.assertEqual([], [r for r in data['issues'] if r['severity']=='BLOCKING'])

    def activate(self,data):
        return self.controller._activate(ProjectDocument.create('我的工程','ui',payload(data)),unsaved=True)

    def test_same_import_input_generated_and_demo_not_production(self):
        path=input_book(self.data,self.root/'arbitrary.xlsx')
        before=hash_file(path)
        self.controller.new_project('保留名称',vnext=True)
        project=self.controller.import_excel(path)
        self.assertEqual('保留名称',project.name)
        data=project.config['vnext']
        self.assertEqual('input',data['provenance']['import_format'])
        self.assertEqual([MODULE,'demo_phy_io_se'],[m['name'] for m in data['modules']])
        self.assertEqual([30,34,40],data['modules'][0]['rons'])
        self.assertIn('IBIS_PIN_MAPPING',self.codes(data))
        target=self.controller.export_config(self.root/'anything.xlsx')
        self.assertEqual('generated',detect(target))
        self.controller.import_excel(target)
        self.assertEqual('generated',self.controller.active_config['vnext']['provenance']['import_format'])
        self.assertEqual(before,hash_file(path))
        demo=read(PACKAGE/'docs/vnext/IBIS_Config_Generated_Example_Final.xlsx')
        self.assertEqual([],demo['modules'])
        self.assertIn('notice',demo['provenance'])
        # The blank Input must never fall through to its populated Demo sheet.
        blank=read(PACKAGE/'docs/vnext/IBIS_Config_Input_Final.xlsx')
        self.assertEqual([],blank['modules'])

    def test_incomplete_export_save_reload_backend_gate_and_isolation(self):
        self.data['voltages']=[]
        data=self.checked()
        self.assertTrue(all(not row[c] for row in data['voltages'] for c in CORNERS))
        self.assertIn('VOLTAGE_MISSING',self.codes(data))
        p=self.activate(data)
        self.controller.save_configuration()
        output=self.controller.export_config(self.root/'incomplete.xlsx')
        reloaded=self.controller.store.load(p.project_id)
        self.assertEqual(data['pins'],reloaded.config['vnext']['pins'])
        with self.assertRaises(ProjectValidationError):
            generate(self.controller.store,self.controller.resources,reloaded)
        with self.assertRaises(ProjectValidationError):
            self.controller.generate_cases()
        self.assertTrue(output.is_file())
        second=self.controller.new_project('另一个项目',vnext=True)
        self.controller.save_configuration()
        self.assertNotEqual(p.project_id,second.project_id)
        self.assertEqual([],second.config['vnext']['modules'])
        self.controller.load_project(p.project_id)
        self.assertEqual(2,len(self.controller.active_config['vnext']['modules']))

    def test_manual_pins_full_order_all_cases_and_sources_readonly(self):
        before={str(p):hash_file(p) for p in (self.root/'sources').rglob('*') if p.is_file()}
        data=self.checked()
        pins=[r for r in data['pins'] if r['module']==MODULE and r['direction']=='DRV']
        self.assertEqual('',next(r for r in pins if r['t2b_pin']=='ipadt')['spf_pin'])
        self.assertEqual('UNRESOLVED',next(r for r in pins if r['t2b_pin']=='ipadt')['status'])
        self.assertEqual('AUTO_MATCH',next(r for r in pins if r['t2b_pin']=='txoe')['status'])
        self.assertEqual(6,len(pins))
        data=confirmed(self.data)
        self.assertReady(data)
        self.assertEqual({'DRV','RCV'},set(d for m,d in core_config(data).templates if m==MODULE))
        self.assertEqual('drv_se',core_config(data).templates[('demo_phy_io_se','DRV')])
        self.activate(data)
        self.controller.save_configuration()
        result=self.controller.generate_cases()
        self.assertTrue(result.ready,repr(result))
        self.assertEqual(8,len(result.cases))
        for case in result.cases:
            self.assertEqual('UNVALIDATED',case.status)
            root=Path(case.path)
            self.assertEqual(7,len(list(root.iterdir())))
            module,direction,_=case.key.split('/')
            top,original=parse_spf(Path(data['resolved']['spf'][module]['files']['TYP']))
            lines=(root/'ckt_topology.sp').read_text().splitlines()
            connection=' '.join(line[3:] if line.startswith('x1 ') else line[2:] for line in lines if line.startswith(('x1 ','+ '))).split()
            self.assertEqual(top,connection[-1])
            aliases={r['spf_pin']:r['t2b_pin'] for r in data['pins'] if r['module']==module and r['direction']==direction}
            self.assertEqual([aliases.get(p,p) for p in original],connection[:-1])
            self.assertFalse((root/'C_comp_view').exists())
            corner=(root/('corner_%s_ron%s_typ.sp'%(direction.lower(),case.key.rsplit('/',1)[1][:-3]))).read_text()
            self.assertIn('.lib ',corner)
            self.assertIn('.inc ',corner)
            self.assertIn('minimum valid code',corner)
        self.assertEqual(before,{str(p):hash_file(p) for p in (self.root/'sources').rglob('*') if p.is_file()})
        self.assertTrue((self.controller.store.project_dir(self.controller.project.project_id)/'logs/generation.json').is_file())

    def test_generated_excel_manual_edits_drive_actual_output(self):
        data=confirmed(self.data)
        target=write(data,self.root/'complete.xlsx')
        book=load_workbook(target)
        sheet=book['Config_Generated']
        for row in sheet:
            if row[0].value=='Voltage' and row[1].value=='VDDQ': row[3].value='0.92'
            if row[0].value=='Timing' and row[1].value=='Tr': row[3].value='31p'
        # Insert a real per-corner signal override, not a hidden metadata change.
        start=next(row[0].row for row in sheet if row[0].value=='06 Signal Override')+2
        sheet.insert_rows(start)
        for col,value in enumerate([MODULE,'DRV','txzqcal_pd[7:0]','','8\'d76',''],1): sheet.cell(start,col,value)
        book.save(target);book.close()
        project=self.controller.import_excel(target)
        self.assertReady(project.config['vnext'])
        self.controller.save_configuration()
        result=self.controller.generate_cases(selected=[(MODULE,'DRV',30)])
        self.assertTrue(result.ready)
        directory=Path(result.cases[0].path)
        text=next(directory.glob('*.t2b')).read_text()
        self.assertIn('[Voltage Range] 0.92 0.72 0.88',text)
        self.assertIn('[Tr] 31p 20p 20p',text)
        corner=(directory/'corner_drv_ron30_typ.sp').read_text()
        self.assertIn("MSI default txzqcal_pd[7:0] = 8'd76",corner)
        self.assertIn('vtxzqcal_pd6 txzqcal_pd[6] 0 dc vdd_typ',corner)
        self.assertIn('vtxzqcal_pd7 txzqcal_pd[7] 0 dc 0',corner)
        with self.assertRaises(ProjectValidationError): self.controller.export_config(target)

    def test_source_changes_stale_manual_pin_msi_and_confirmation_persists(self):
        for kind in ('spf','template','msi'):
            with self.subTest(kind=kind):
                data=confirmed(self.data)
                target=write(data,self.root/(kind+'.xlsx'))
                if kind=='spf': path=Path(data['resolved']['spf'][MODULE]['files']['TYP'])
                elif kind=='template': path=Path(data['resolved']['templates'][MODULE+'|DRV']['path'])
                else: path=Path(data['msi_mappings'][0]['file'])
                original=path.read_bytes()
                if kind=='msi':
                    book=load_workbook(path);book[MODULE].cell(2,9,'source changed');book.save(path);book.close()
                else: path.write_bytes(original+b'\n* changed\n')
                checked=precheck(read(target),TEMPLATES)
                manual=next(r for r in checked['pins'] if r['module']==MODULE and r['direction']=='DRV' and r['t2b_pin']=='ipadt')
                self.assertEqual('STALE',manual['status'])
                self.assertEqual('IOPADT',manual['spf_pin'])
                self.assertIn('IBIS_PIN_MAPPING',self.codes(checked))
                for field in ('pins','msi_mappings'):
                    for row in checked[field]:
                        if row['status']=='STALE': row.update(status='MANUAL',fingerprint=row['current_fingerprint'])
                checked=precheck(checked,TEMPLATES)
                self.assertReady(checked)
                write(checked,target)
                self.assertReady(precheck(read(target),TEMPLATES))
                path.write_bytes(original)

    def test_missing_ambiguous_and_inconsistent_inputs_located(self):
        data=confirmed(self.data)
        source=Path(data['resolved']['spf'][MODULE]['files']['TYP'])
        duplicate=source.with_name('extra_'+source.name)
        duplicate.write_bytes(source.read_bytes())
        self.assertIn('SPF_MATCH',self.codes(precheck(data,TEMPLATES)))
        duplicate.unlink()
        text=source.read_text();source.write_text(text.replace('TXOE','TXOE_CHANGED'))
        checked=precheck(data,TEMPLATES)
        self.assertIn('SPF_INTERFACE',self.codes(checked))
        self.assertTrue(any(r['location']==str(source) and r['corner']=='TYP' for r in checked['issues']))
        source.write_text(text)
        source.unlink()
        self.assertIn('SPF_MATCH',self.codes(precheck(data,TEMPLATES)))
        data['msi_mappings']=[]
        msi=Path(data['sources']['msi'])/'demo_msi.xlsx'
        shutil.copyfile(msi,msi.with_name('duplicate.xlsx'))
        checked=precheck(data,TEMPLATES)
        self.assertIn('MSI_MAPPING',self.codes(checked))
        self.assertEqual('',checked['msi_mappings'][0]['file'])
        self.assertGreaterEqual(len(checked['resolved']['candidates']),2)

    def test_anchors_single_template_scope_process_sections(self):
        data=confirmed(self.data)
        data['sources']['model']=str(Path(data['sources']['model'])/'rbd.lib')
        data['sources']['msi']=str(Path(data['sources']['msi'])/'demo_msi.xlsx')
        data['sources']['spf']=data['resolved']['spf'][MODULE]['files']['TYP']
        self.assertReady(precheck(data,TEMPLATES))
        data['sources']['template']=data['resolved']['templates'][MODULE+'|DRV']['path']
        self.assertIn('TEMPLATE_SCOPE',self.codes(precheck(data,TEMPLATES)))
        data['sources']['template']=''
        data['corners']['MAX']['process']='FS'
        data['corners']['MIN']['process']='SF'
        checked=precheck(data,TEMPLATES)
        # New template source makes earlier manual mapping stale: reconfirm.
        checked=confirmed(checked)
        for row in checked['pins']: row.update(status='MANUAL',fingerprint=row['current_fingerprint'])
        checked=precheck(checked,TEMPLATES)
        self.assertReady(checked)
        self.assertEqual(('rbd_fs','cap_fs','mos_var_fs'),core_config(checked).process_sections['spectre:FS'])
        path=Path(data['sources']['model']).parent/'cap.lib'
        path.write_text('.lib cap_f\n.endl cap_f\n')
        self.assertIn('PDK_SECTION',self.codes(precheck(data,TEMPLATES)))

    def test_expanded_input_rows_ron_values_and_compact_export(self):
        self.assertEqual([30,34,40],parse_rons('30，34; 40 30'))
        for value in ('0','-1','30.5',''):
            with self.assertRaises(ValueError): parse_rons(value)
        for i in range(4):
            self.data['modules'].append(dict(self.data['modules'][0],name='new_module_'+str(i)))
        self.data['voltages'].extend(dict(domain='EXTRA'+str(i),MAX='1',TYP='1',MIN='1') for i in range(4))
        self.data['signals']=[dict(signal='signal_'+str(i),MAX='',TYP='0',MIN='') for i in range(4)]
        path=input_book(self.data,self.root/'inserted.xlsx')
        data=read(path)
        self.assertEqual(6,len(data['modules']))
        self.assertEqual(7,len(data['voltages']))
        self.assertEqual(4,len(data['signals']))
        data=confirmed(fixture(self.root/'second'))
        exported=write(data,self.root/'compact.xlsx')
        book=load_workbook(exported)
        sheet=book['Config_Generated']
        self.assertLess(sheet.max_row,70)
        for row in sheet:
            first=str(row[0].value or '')
            if first in ('Voltage','Timing',MODULE):
                self.assertTrue(all(c.fill.patternType is None for c in row))
                self.assertTrue(all(c.border.left.style=='thin' for c in row))
        book.close()
        self.assertEqual(2,len(data['appends']))
        self.assertEqual(2,len(data['overrides']))

    def test_logic_formats_global_module_inheritance_conflicts_width(self):
        for raw in ("8'h4C","8'b01001100","8'd76",'0x4C','0b01001100','76'):
            self.assertEqual('01001100',parse_logic(raw,MODULE,'test',8))
        for raw in ("8'b2","8'd256","4'h1F",'256','-1','x',''):
            with self.assertRaises(Exception): parse_logic(raw,MODULE,'test',8)
        data=confirmed(self.data)
        data['signals']=[dict(signal='txzqcal_pd[7:0]',MAX='0x4C',TYP='',MIN='')]
        data['overrides']=[dict(module=MODULE,direction='DRV',signal='txzqcal_pd[7:0]',MAX="8'd12",TYP='',MIN='')]
        checked=precheck(data,TEMPLATES)
        self.assertReady(checked)
        config=core_config(checked)
        self.assertEqual("8'd12",next(r['value'] for r in config.overrides if r['module']==MODULE and r['direction']=='DRV'))
        self.assertEqual('0x4C',next(r['value'] for r in config.overrides if r['module']=='demo_phy_io_se'))
        data['overrides'].append(dict(data['overrides'][0]))
        self.assertIn('SIGNAL_CONFLICT',self.codes(precheck(data,TEMPLATES)))
        data['overrides']=[dict(module=MODULE,direction='DRV',signal='reg_txslice_pd[3:0]',MAX='0',TYP='',MIN='')]
        self.assertIn('SIGNAL_VALUE',self.codes(precheck(data,TEMPLATES)))
        path=Path(data['sources']['msi'])/'demo_msi.xlsx'
        book=load_workbook(path);book[MODULE].cell(9,2,3);book.save(path);book.close()
        self.assertIn('MSI_MAPPING',self.codes(precheck(data,TEMPLATES)))

    def test_reimport_minimal_preserves_manual_and_orphans_visible(self):
        data=confirmed(self.data)
        self.activate(data)
        path=input_book(self.data,self.root/'input.xlsx')
        self.controller.import_excel(path)
        checked=self.controller.active_config['vnext']
        self.assertReady(checked)
        self.assertEqual('MANUAL',next(r['status'] for r in checked['pins'] if r['t2b_pin']=='ipadt'))
        checked['modules']=checked['modules'][1:]
        orphan=precheck(checked,TEMPLATES)
        self.assertIn('PIN_ORPHAN',self.codes(orphan))
        self.assertTrue(any(r['module']==MODULE and r['status']=='STALE' for r in orphan['pins']))

    def test_cli_preprocess_exports_blocked_configuration(self):
        path=input_book(self.data,self.root/'input.xlsx')
        target=self.root/'cli.xlsx'
        result=subprocess.run([sys.executable,'-m','ibis_ui.vnext','preprocess','--input',str(path),'--workspace',str(self.root/'cli-work'),'--output',str(target)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,universal_newlines=True)
        self.assertEqual(0,result.returncode,result.stdout+result.stderr)
        self.assertTrue(target.is_file())
        self.assertIn('阻断项',result.stdout)


    def test_conflicting_mapping_rows_survive_repeated_precheck(self):
        for field,code in [('pins','PIN_DUPLICATE'),('msi_mappings','MSI_DUPLICATE')]:
            data=confirmed(self.data)
            original_count=len(data[field])
            data[field].append(copy.deepcopy(data[field][0]))
            for _ in range(3):
                data=precheck(data,TEMPLATES)
                self.assertIn(code,self.codes(data))
                self.assertEqual(original_count+1,len(data[field]))
                self.assertTrue(write(data,self.root/'conflict.xlsx').is_file())
            data[field].pop()
            self.assertReady(precheck(data,TEMPLATES))

    def test_blocking_validation_matrix_and_untrusted_status(self):
        ready=confirmed(self.data)
        mutations=[
            ('MODULE_INVALID',lambda d:d['modules'][0].update(ron='-4')),
            ('MODULE_INVALID',lambda d:d['modules'][0].update(topology='AUTO')),
            ('CORNER_INVALID',lambda d:d['corners']['TYP'].update(tr='0p')),
            ('CORNER_INVALID',lambda d:d['corners']['TYP'].update(tf='nan')),
            ('CORNER_INVALID',lambda d:d['corners']['TYP'].update(temperature='100')),
            ('VOLTAGE_MISSING',lambda d:d['voltages'][0].update(TYP='NaN')),
            ('SPF_REUSED',lambda d:d['corners']['MAX'].update(spf_key='typical')),
            ('SOURCE_INVALID',lambda d:d['sources'].update(msi=str(self.root/'missing.xlsx'))),
            ('IBIS_PIN_MAPPING',lambda d:next(r for r in d['pins'] if r['t2b_pin']=='ipadt').update(status='AUTO_MATCH')),
            ('MSI_UNCONFIRMED',lambda d:d['msi_mappings'][0].update(status='UNRESOLVED')),
        ]
        for code,change in mutations:
            with self.subTest(code=code):
                data=copy.deepcopy(ready);change(data)
                checked=precheck(data,TEMPLATES)
                self.assertIn(code,self.codes(checked))
                from ibis_ui.config_adapter import payload_to_core_config
                with self.assertRaises(ProjectValidationError): payload_to_core_config(payload(checked))
                self.assertTrue(write(checked,self.root/'blocked.xlsx').is_file())

    def test_selected_single_template_anchor_and_verified_fs_generate(self):
        data=copy.deepcopy(self.data)
        data['modules']=data['modules'][:1]
        data['modules'][0]['generate_type']='DRV'
        data['sources']['template']=str(Path(data['sources']['template'])/'drv_diff.t2b')
        data['corners']['MAX']['process']='FS'
        data['corners']['MIN']['process']='SF'
        checked=confirmed(data)
        self.assertReady(checked)
        self.activate(checked);self.controller.save_configuration()
        result=self.controller.generate_cases(selected=[(MODULE,'DRV',30)])
        self.assertTrue(result.ready)
        case=Path(result.cases[0].path)
        self.assertIn('rbd_fs',(case/'corner_drv_ron30_max.sp').read_text())
        self.assertIn('mos_var_sf',(case/'corner_drv_ron30_min.sp').read_text())

    def test_large_metadata_roundtrip_and_literal_excel_strings(self):
        data=confirmed(self.data)
        data['provenance']['long_note']='人工说明'*12000
        data['appends'][0]['body']='=literal SPICE note'
        path=write(data,self.root/'large.xlsx')
        loaded=read(path)
        self.assertEqual(data['provenance']['long_note'],loaded['provenance']['long_note'])
        self.assertEqual('=literal SPICE note',loaded['appends'][0]['body'])
        self.assertEqual(data['pins'][0]['fingerprint'],loaded['pins'][0]['fingerprint'])

    def test_manual_excel_pin_edit_and_readonly_reimported_autosave(self):
        data=self.checked()
        path=write(data,self.root/'mapping.xlsx')
        book=load_workbook(path)
        for row in book['Config_Generated']:
            if row[2].value in ('ipadt','ipadc','ipad') and row[0].value in (MODULE,'demo_phy_io_se'):
                row[3].value={'ipadt':'IOPADT','ipadc':'IOPADC','ipad':'IOPAD'}[row[2].value]
        book.save(path);book.close()
        self.controller.import_excel(path)
        self.assertReady(self.controller.active_config['vnext'])
        self.controller.save_configuration()
        project_path=self.controller.store.project_dir(self.controller.project.project_id)
        generated=project_path/'config/Config_Generated.xlsx'
        checksum=hash_file(generated)
        self.controller.import_excel(generated)
        self.controller.draft['vnext']['corners']['TYP']['tr']='40p'
        self.controller.save_configuration()
        self.assertEqual(checksum,hash_file(generated))
        self.assertTrue((project_path/'config/Config_Current.xlsx').exists())

    def test_workspace_cannot_write_inside_input_and_source_status_tamper(self):
        data=confirmed(self.data)
        store=ProjectStore(Path(data['sources']['spf'])/'workspace')
        project=ProjectDocument.create('禁止污染输入','ui',payload(data))
        from ibis_ui.vnext.service import save_checked
        with self.assertRaises(ProjectValidationError): save_checked(store,project)
        # A false green status persisted by an external tool cannot bypass backend.
        data['issues']=[]
        data['corners']['TYP']['tr']=''
        with self.assertRaises(ProjectValidationError):
            generate(self.controller.store,self.controller.resources,ProjectDocument.create('检查','ui',payload(data)))


@unittest.skipUnless(os.name == 'nt' or os.environ.get('DISPLAY'), '真实 Tk 测试需要 DISPLAY')
class VNextGuiTests(unittest.TestCase):
    def test_ui_edit_issue_location_export_restore_and_generate(self):
        from ibis_ui.view import IBISApplicationView
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary)
            data=fixture(root)
            input_path=input_book(data,root/'input.xlsx')
            controller=ProjectController(ProjectStore(root/'workspace'))
            controller.new_project('GUI工程',vnext=True)
            controller.import_excel(input_path)
            app=IBISApplicationView(controller)
            try:
                app._show_main();app.update()
                panel=app.vnext_panel
                self.assertTrue(panel.winfo_ismapped())
                self.assertEqual('disabled',str(panel.generate_button['state']))
                issue=next(i for i,r in enumerate(panel.data['issues']) if r['code']=='IBIS_PIN_MAPPING')
                panel.issue_tree.selection_set(str(issue));app.update()
                self.assertEqual(str(panel.config),panel.notebook.select())
                panel.flush()
                for i,row in enumerate(list(panel.data['pins'])):
                    if row['t2b_pin'] in ('ipadt','ipadc'):
                        panel.editors['pins'].set_value(i,'spf_pin','IOPADT' if row['t2b_pin']=='ipadt' else 'IOPADC')
                panel.editors['corners'].set_value(2,'TYP','35p')
                panel.check();app.update()
                self.assertEqual('normal',str(panel.generate_button['state']))
                controller.save_configuration()
                out=controller.export_config(root/'gui.xlsx')
                self.assertEqual('35p',read(out)['corners']['TYP']['tr'])
                project_id=controller.project.project_id
                controller.close_project('save');controller.load_project(project_id)
                app._show_main();app.update()
                panel=app.vnext_panel
                self.assertEqual('35p',panel.data['corners']['TYP']['tr'])
                panel.generate()
                import time
                deadline=time.monotonic()+20
                while panel.busy and time.monotonic()<deadline:
                    app.update();time.sleep(.01)
                self.assertFalse(panel.busy)
                self.assertEqual(6,len(controller.generation_statuses()))
                self.assertTrue(all(r['status']=='UNVALIDATED' for r in controller.generation_statuses()))
            finally:
                app.destroy()
