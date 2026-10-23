from __future__ import annotations
import copy
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from ibis_ui.controller import ProjectController
from ibis_ui.project_store import ProjectStore
from ibis_ui.runtime import ResourceResolver
from ibis_ui.effective_msi import inspect, prepare, file_sha256, workbook_diff
from ibis_ui.configuration_review import changes, append_diagnostics
from ibis_ui.errors import ProjectValidationError, ProjectStorageError

ROOT = Path(__file__).resolve().parents[1]
MODULE = 'demo_phy_io_diff'
KEY = (MODULE, 'DRV', 40)


class UI042Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name).resolve()
        self.c = ProjectController(ProjectStore(self.root / 'workspace'), resources=ResourceResolver(ROOT))
        self.c.import_excel(ROOT / 'inputs/config/ibis_config.xlsx')
        shutil.copytree(ROOT / 'inputs/msi', self.root / 'msi')
        self.c.update_root('msi', str(self.root / 'msi'))
        self.c.update_module(MODULE, 'generate_type', 'BOTH')
        self.c.update_module(MODULE, 'impedances', [40])
        self.c.save_project_draft()
        self.raw = self.root / 'msi/demo_msi.xlsx'
        self.raw_hash = file_sha256(self.raw)

    def test_master_is_actual_core_input_and_output_snapshot(self):
        from ibis_auto.service import generate_from_config
        seen = []
        def capture(config, *args, **kwargs):
            seen.append(copy.deepcopy(config))
            return generate_from_config(config, *args, **kwargs)
        with patch('ibis_ui.generation_service.generate_from_config', side_effect=capture):
            result = self.c.generate_cases()
        self.assertTrue(result.ready, [r.message for r in result.cases])
        self.assertEqual({r.status for r in result.cases}, {'UNVALIDATED'})
        config = seen[0]
        self.assertIn('_inputs', config.roots['msi'].parts)
        self.assertEqual(config.msi_mappings[MODULE][0], 'MSI_Master.xlsx')
        master = config.roots['msi'] / 'MSI_Master.xlsx'
        for record in self.c.project.workflow['generation']['cases'].values():
            manifest = json.loads(Path(record['manifest_path']).read_text(encoding="utf-8"))
            self.assertEqual(file_sha256(master), manifest['master_sha256'])
            self.assertEqual(str(master), manifest['master_path'])
        self.assertEqual(file_sha256(self.raw), self.raw_hash)
        self.assertTrue((master.parent / 'source_manifest.json').is_file())
        self.assertEqual(len(list((master.parents[1]).rglob('MSI_Master.xlsx'))), 1)

    def test_attribute_materialized_but_corner_override_kept_as_case_rule(self):
        c = self.c
        c.replace_config_rows('attribute_overrides', [dict(module=MODULE, signal='reg_txffe[3:0]', attribute='POWER_DOMAIN', value='VDDQ')])
        c.replace_config_rows('overrides', [dict(module=MODULE, direction='RCV', corner='TYP', signal='reg_txffe[3:0]', value="4'b0001")])
        c.save_project_draft()
        result = c.generate_cases()
        self.assertTrue(result.ready, [r.message for r in result.cases])
        state = c.project.workflow['msi_master']
        book = load_workbook(state['path'], read_only=True)
        sheet = book[state['sources'][0]['master_sheet']]
        rows = list(sheet.values)
        headers = list(rows[0])
        from ibis_auto.msi import _columns
        cols = _columns(headers)
        row = next(r for r in rows[1:] if r[cols['signal_name']] == 'reg_txffe[3:0]')
        self.assertEqual(row[cols['power_domain']], 'VDDQ')
        self.assertEqual(row[cols['default']], "4'b0000")
        self.assertIn('_Case_Overrides', book.sheetnames)
        book.close()
        for case in result.cases:
            files = list(Path(case.path).glob('corner*.sp'))
            for file in files:
                line = next(line for line in file.read_text(encoding="utf-8").splitlines() if line.startswith('vreg_txffe0 '))
                self.assertIn('vddq_typ' if '/RCV/' in case.key and '_typ.sp' in file.name else 'dc 0', line)

    def test_options_append_alias_and_t2b_reach_generated_files(self):
        c = self.c
        options = ['.option method=gear', '.option runlvl=6', '.option reltol=1e-5']
        c.replace_config_rows('simulation_options', options)
        body = '* 多行追加\nRextra ipadt local_node 100\nCextra local_node 0 1p\n'
        c.replace_config_rows('appends', [dict(module=MODULE, direction='DRV', body=body)])
        aliases = copy.deepcopy(c.active_config['aliases'])
        aliases[0]['alias'] = 'changed_alias'
        c.replace_config_rows('aliases', aliases)
        c.update_module(MODULE, 'tr.TYP', '31p')
        c.update_corner_voltage('VDDQ', 'TYP', '0.54')
        c.save_project_draft()
        preview = c.preview_t2b(KEY)
        result = c.generate_cases([KEY])
        self.assertTrue(result.ready, result.cases)
        output = Path(result.cases[0].path)
        self.assertEqual(next(output.glob('*.t2b')).read_text(encoding="utf-8"), preview)
        self.assertIn('31p', preview)
        self.assertIn('0.54', preview)
        self.assertIn(body, (output / 'ckt_topology.sp').read_text(encoding="utf-8"))
        self.assertIn('changed_alias', (output / 'ckt_topology.sp').read_text(encoding="utf-8"))
        text = (output / 'corner_drv_ron40_typ.sp').read_text(encoding="utf-8")
        self.assertLess(text.index(options[0]), text.index(options[1]))
        self.assertLess(text.index(options[1]), text.index(options[2]))

    def test_source_change_rebuilds_without_destroying_old_snapshot(self):
        c = self.c
        self.assertTrue(c.generate_cases([KEY]).ready)
        old = copy.deepcopy(c.project.workflow['msi_master'])
        book = load_workbook(self.raw)
        book.worksheets[0].cell(2, 11, 'source updated')
        book.save(self.raw)
        book.close()
        self.assertEqual(c.generation_statuses()[0]['status'], 'STALE')
        self.assertTrue(c.effective_msi_preview().issues)
        self.assertTrue(c.generate_cases([KEY]).ready)
        new = c.project.workflow['msi_master']
        self.assertNotEqual(old['sha256'], new['sha256'])
        self.assertEqual(file_sha256(Path(old['path'])), old['sha256'])
        self.assertEqual(len(new['history']), 2)

    def test_repeated_generation_reuses_identical_master(self):
        self.c.generate_cases([KEY])
        first = copy.deepcopy(self.c.project.workflow['msi_master'])
        self.c.generate_cases([KEY])
        second = self.c.project.workflow['msi_master']
        self.assertEqual(first['path'], second['path'])
        self.assertEqual(first['sha256'], second['sha256'])
        self.assertEqual(len(second['history']), 1)

    def test_manual_master_edit_blocks_until_explicit_restore(self):
        self.c.generate_cases([KEY])
        state = self.c.project.workflow['msi_master']
        path = Path(state['path'])
        book = load_workbook(path)
        book.worksheets[0].cell(2, 11, 'manual change')
        book.save(path)
        book.close()
        manual_hash = file_sha256(path)
        result = self.c.generate_cases([KEY])
        self.assertFalse(result.ready)
        self.assertIn('MSI_MASTER_MANUAL_EDIT', result.cases[0].message)
        self.assertEqual(file_sha256(path), manual_hash)
        restored = self.c.prepare_master(restore=True)
        self.assertNotEqual(restored, path)
        backups = list((self.c.store.project_dir(self.c.project.project_id) / 'msi/manual_backups').rglob('*.xlsx'))
        self.assertEqual(file_sha256(backups[0]), manual_hash)
        self.c.save_project_draft()
        self.assertTrue(self.c.generate_cases([KEY]).ready)

    def test_mapping_change_uses_new_workbook_and_sheet(self):
        self.c.generate_cases([KEY])
        second = self.root / 'msi/alternate.xlsx'
        shutil.copyfile(self.raw, second)
        book = load_workbook(second)
        book.worksheets[0].title = 'new_sheet'
        book.save(second)
        book.close()
        self.c.replace_config_rows('msi_mappings', [dict(module=MODULE, file=str(second), sheet='new_sheet', column='PIN_NAME')])
        # Fixture header may be an alias; retain the actual header.
        from ibis_auto.ooxml import read_workbook
        from ibis_auto.msi import _columns
        header = read_workbook(second)['new_sheet'][0]
        self.c.draft['msi_mappings'][0]['column'] = header[_columns(header)['signal_name']]
        self.c.save_project_draft()
        self.assertTrue(self.c.generate_cases([KEY]).ready)
        source = self.c.project.workflow['msi_master']['sources'][0]
        self.assertEqual(source['source_workbook'], str(second))
        self.assertEqual(source['source_sheet'], 'new_sheet')

    def test_mapping_errors_missing_sheet_column_duplicate_width(self):
        from ibis_auto.ooxml import read_workbook
        from ibis_auto.msi import _columns
        sheets = read_workbook(self.raw)
        sheet = next(s for s in sheets if MODULE in s)
        header = sheets[sheet][0]
        column = header[_columns(header)['signal_name']]
        for bad_sheet, bad_col in [('missing', column), (sheet, 'missing')]:
            with self.subTest(sheet=bad_sheet, column=bad_col):
                self.c.draft['msi_mappings'] = [dict(module=MODULE, file=str(self.raw), sheet=bad_sheet, column=bad_col)]
                report = self.c.effective_msi_preview()
                self.assertTrue(report.errors)
        self.c.draft['msi_mappings'] = []
        book = load_workbook(self.raw)
        ws = book[sheet]
        ws.append([cell.value for cell in ws[2]])
        book.save(self.raw)
        book.close()
        self.assertIn('重复', self.c.effective_msi_preview().errors[0]['message'])
        self.assertTrue(self.c.effective_msi_preview().errors[0]['signal'])

    def test_alias_duplicate_invalid_reference_and_append_duplicate(self):
        before = self.c.project.to_dict()
        for field, rows in [('aliases', [dict(module=MODULE, direction='DRV', original_pin='VDD', alias='x'), dict(module=MODULE, direction='BOTH', original_pin='VSS', alias='x')]), ('aliases', [dict(module=MODULE, direction='DRV', original_pin='NOT_A_PIN', alias='x')]), ('appends', [dict(module=MODULE, direction='DRV', body='Rdup a b 1\nRdup c d 2')])]:
            self.c.discard_draft()
            self.c.replace_config_rows(field, rows)
            with self.assertRaises(ProjectValidationError):
                self.c.save_project_draft()
            self.assertEqual(before, self.c.project.to_dict())
        self.assertTrue(all(level == 'Warning' for level, _ in append_diagnostics('Xsub internal 0 complex_model', ['VDD'])))

    def test_save_diff_tracks_saved_not_applied_and_multiline_changes(self):
        self.c.replace_config_rows('appends', [dict(module=MODULE, direction='DRV', body='Rone a 0 1\nRtwo b 0 2')])
        self.c.save_project_draft()
        self.c.draft['appends'][0]['body'] = 'Rone a 0 9\nRthree c 0 3'
        self.c.apply_changes()
        rows = self.c.review_changes()
        row = next(r for r in rows if r['field'].endswith('body'))
        self.assertIn('-Rtwo', row['detail'])
        self.assertIn('+Rthree', row['detail'])
        self.assertEqual(row['scope'], MODULE + ' / DRV')
        before = self.c.store.load(self.c.project.project_id).to_dict()
        with patch.object(self.c.store, 'save', side_effect=ProjectStorageError('TEST', '测试保存失败')):
            with self.assertRaises(ProjectStorageError):
                self.c.save_project_draft()
        self.assertEqual(before, self.c.store.load(self.c.project.project_id).to_dict())
        self.assertIn('Rthree', self.c.draft['appends'][0]['body'])

    def test_external_template_preview_generation_and_changed_notice(self):
        path = self.root / 'external.t2b'
        source = self.c.template_source(MODULE, 'DRV')
        path.write_text(source.replace('[Source] Akrostar', '[Source] External'), encoding='utf-8')
        self.c.set_template_source(MODULE, 'DRV', '外部模板', str(path))
        self.c.save_project_draft()
        self.assertIn('External', self.c.preview_t2b(KEY))
        self.assertTrue(self.c.generate_cases([KEY]).ready)
        path.write_text(source.replace('[Source] Akrostar', '[Source] Changed'), encoding='utf-8')
        self.assertTrue(self.c.template_origin(MODULE, 'DRV')['changed'])
        self.assertIn('Changed', self.c.preview_t2b(KEY))
        self.c.set_template_source(MODULE, 'DRV', '内置母模板')
        self.assertNotIn('[Source] Changed', self.c.preview_t2b(KEY))

    def test_case_override_conflict_and_invalid_signal_rejected(self):
        for rows in [[dict(module=MODULE, direction='DRV', corner='ALL', signal='reg_txffe[3:0]', value='0'), dict(module=MODULE, direction='BOTH', corner='TYP', signal='reg_txffe[3:0]', value='1')], [dict(module=MODULE, direction='DRV', corner='TYP', signal='bad_signal', value='1')]]:
            self.c.replace_config_rows('overrides', rows)
            with self.assertRaises(ProjectValidationError):
                self.c.save_project_draft()

    def test_cancel_does_not_prepare_master_or_alter_config(self):
        before = copy.deepcopy(self.c.project.config)
        result = self.c.generate_cases([KEY], cancel_requested=lambda: True)
        self.assertEqual(result.cases[0].status, '已取消')
        self.assertEqual(before, self.c.project.config)
        self.assertFalse(self.c.project.workflow['msi_master']['path'])

    def test_modified_output_master_invalidates_check(self):
        self.c.generate_cases([KEY])
        record = self.c.project.workflow['generation']['cases'][f'{MODULE}/DRV/40ohm']
        manifest = json.loads(Path(record['manifest_path']).read_text(encoding="utf-8"))
        Path(manifest['master_path']).write_bytes(b'broken')
        status = next(r for r in self.c.generation_statuses() if r['direction'] == 'DRV')
        self.assertEqual(status['status'], 'FAIL')
        self.assertIn('Master', status['message'])

    def test_multiple_workbooks_sheets_and_per_module_attributes(self):
        other = 'other_lane_diff'
        self.c.duplicate_module(MODULE, other)
        second = self.root / 'msi/second.xlsx'
        book = load_workbook(self.raw)
        original_sheet = book.worksheets[0].title
        book.copy_worksheet(book.worksheets[0]).title = 'second_sheet'
        book.save(second)
        book.close()
        from ibis_auto.ooxml import read_workbook
        from ibis_auto.msi import _columns
        header = read_workbook(self.raw)[original_sheet][0]
        column = header[_columns(header)['signal_name']]
        self.c.draft['msi_mappings'] = [dict(module=MODULE, file=str(self.raw), sheet=original_sheet, column=column), dict(module=other, file=str(second), sheet='second_sheet', column=column)]
        self.c.draft['attribute_overrides'] = [dict(module=other, signal='reg_txffe[3:0]', attribute='POWER_DOMAIN', value='VDDQ')]
        report = self.c.effective_msi_preview()
        report.require_valid()
        self.assertEqual(len(report.sheets), 2)
        self.assertEqual(len({r['source_workbook'] for r in report.mappings}), 2)
        one = next(r for r in report.signals[MODULE] if r['signal'] == 'reg_txffe[3:0]')
        two = next(r for r in report.signals[other] if r['signal'] == 'reg_txffe[3:0]')
        self.assertNotEqual(one['effective_power_domain'], two['effective_power_domain'])
        # Both modules can also map different sheets of one workbook.
        self.c.draft['msi_mappings'][0]['file'] = str(second)
        report = self.c.effective_msi_preview()
        report.require_valid()
        self.assertEqual(len({r['source_workbook'] for r in report.mappings}), 1)
        self.assertEqual(len(report.sheets), 2)

    def test_invalid_width_attribute_conflict_and_legacy_master_recovery(self):
        for overrides in [[dict(module=MODULE, signal='missing', attribute='POWER_DOMAIN', value='VDD')], [dict(module=MODULE, signal='TXOE', attribute='POWER_DOMAIN', value='VDD')] * 2]:
            self.c.draft['attribute_overrides'] = overrides
            self.assertTrue(self.c.effective_msi_preview().errors)
        self.c.discard_draft()
        # Legacy manual merge API remains readable, but manual Master edits are never silently adopted.
        old = self.c.merge_msi_master()
        self.assertTrue(old.is_file())
        self.c.prepare_master(restore=True)
        self.c.save_project_draft()
        self.assertTrue(self.c.generate_cases([KEY]).ready)
        self.assertEqual(file_sha256(self.raw), self.raw_hash)
        # Invalid width gets a located diagnostic.
        book = load_workbook(self.raw)
        from ibis_auto.msi import _columns
        ws = book.worksheets[0]
        cols = _columns([str(c.value or '') for c in ws[1]])
        ws.cell(2, cols['width'] + 1, 'invalid-width')
        book.save(self.raw)
        book.close()
        self.assertIn('WIDTH', self.c.effective_msi_preview().errors[0]['message'])
        self.assertTrue(self.c.effective_msi_preview().errors[0]['signal'])

    def test_failed_new_case_keeps_previous_success_manifest_and_outputs(self):
        self.c.generate_cases([KEY])
        old = copy.deepcopy(self.c.project.workflow['generation']['cases'][f'{MODULE}/DRV/40ohm'])
        old_manifest = Path(old['manifest_path']).read_bytes()
        from ibis_auto.errors import IbisError
        with patch('ibis_auto.service.generate_case', side_effect=IbisError('TEST_FAILURE', '测试生成失败')):
            result = self.c.generate_cases([KEY])
        self.assertEqual(result.cases[0].status, 'FAIL')
        new = self.c.project.workflow['generation']['cases'][f'{MODULE}/DRV/40ohm']
        self.assertNotIn('manifest_path', new)
        self.assertEqual(Path(old['manifest_path']).read_bytes(), old_manifest)
        self.assertTrue(Path(old['output_path']).is_dir())

    def test_core_input_snapshot_is_immutable_and_file_check_detects_tampering(self):
        result = self.c.generate_cases([KEY])
        output = Path(result.cases[0].path)
        (output / 'component.sp').write_text('tampered', encoding='utf-8')
        status, message = self.c.generation_service().validate_existing_case(self.c.project, KEY)
        self.assertEqual(status, 'FAIL')
        self.assertIn('component.sp', message)

    def test_unrecognized_spice_and_scoped_instance_names_are_warnings_only(self):
        body = '.subckt a p n\nR1 p n 1\n.ends\n.subckt b p n\nR1 p n 2\n.ends\nXunknown a 0 model'
        issues = append_diagnostics(body)
        self.assertFalse(any(level == 'Error' for level, _ in issues))
