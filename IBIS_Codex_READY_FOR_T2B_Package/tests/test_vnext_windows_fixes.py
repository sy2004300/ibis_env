from __future__ import annotations
import copy
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from ibis_auto.values import parse_logic
from ibis_ui.vnext.logic_codes import number, display_row, change_radix, edit_code, canonical_value
from ibis_ui.vnext.excel import read, write
from ibis_ui.vnext.scan import precheck
from ibis_ui.vnext.adapter import core_config
from ibis_ui.vnext.model import CORNERS, payload
from ibis_ui.vnext.service import generate
from ibis_ui.models import ProjectDocument
from ibis_ui.project_store import ProjectStore
from ibis_ui.runtime import ResourceResolver
from vnext_fixture import fixture, confirmed, TEMPLATES, MODULE, input_book

SIGNAL='reg_txffe[3:0]'


class SignalAndAppendTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.data=confirmed(fixture(self.root))

    def tearDown(self): self.temp.cleanup()

    def row(self,radix='二进制',value='1111'):
        return dict(signal=SIGNAL,radix=radix,encoding='code',**{c:value for c in CORNERS})

    def assertReady(self,data):
        self.assertEqual([], [r for r in data['issues'] if r['severity']=='BLOCKING'])

    def test_equivalent_radices_and_generated_six_case_sources(self):
        expected=None
        for radix,value,literal in [('二进制','1111',"4'b1111"),('十六进制','F',"4'hF"),('十进制','15',"4'd15")]:
            data=copy.deepcopy(self.data)
            data['signals']=[self.row(radix,value)]
            checked=precheck(data,TEMPLATES)
            self.assertReady(checked)
            config=core_config(checked)
            self.assertTrue(all(r['value']==literal for r in config.overrides))
            project=ProjectDocument.create(radix,'ui',payload(checked))
            result=generate(ProjectStore(self.root/'workspace'),ResourceResolver.current(),project)
            self.assertTrue(result.ready, str(result))
            self.assertEqual(6,len(result.cases))
            actual=[]
            for case in result.cases:
                path=Path(case.path)
                for corner in sorted(path.glob('corner_*.sp')):
                    lines=[line for line in corner.read_text().splitlines() if line.startswith('vreg_txffe')]
                    self.assertEqual(4,len(lines))
                    self.assertTrue(all('dc vdd_' in line for line in lines))
                    actual.append(lines)
            if expected is None: expected=actual
            self.assertEqual(expected,actual)

    def test_schema1_migration_keeps_meaning_and_excel_version_guard(self):
        from ibis_ui.vnext.model import validate_shape
        from ibis_ui.errors import ExcelImportError
        import json
        old=copy.deepcopy(self.data)
        old['schema_version']=1
        old['signals']=[dict(signal=SIGNAL,MAX='10',TYP='0xF',MIN="4'b0001")]
        validate_shape(old)
        new=precheck(old,TEMPLATES)
        self.assertReady(new)
        self.assertEqual(2,new['schema_version'])
        self.assertEqual('1010',parse_logic(core_config(new).overrides[0]['value'],MODULE,SIGNAL,4))
        path=write(new,self.root/'migration.xlsx')
        loaded=precheck(read(path),TEMPLATES)
        self.assertEqual('1010',parse_logic(core_config(loaded).overrides[0]['value'],MODULE,SIGNAL,4))
        book=load_workbook(path)
        metadata=json.loads(book['_VNext'].cell(1,2).value)
        self.assertEqual(2,metadata['schema_version'])
        metadata['schema_version']=999
        book['_VNext'].cell(1,2,json.dumps(metadata));book.save(path);book.close()
        with self.assertRaises(ExcelImportError): read(path)

    def test_switch_preserves_numbers_and_blanks_and_no_mutation_on_error(self):
        row=self.row();row['MIN']=''
        resolved=self.data['resolved']
        for base,shown in [('十六进制','F'),('十进制','15'),('二进制','1111')]:
            change_radix(row,base,resolved,True)
            self.assertEqual(shown,row['MAX'])
            self.assertEqual('',row['MIN'])
            self.assertEqual('1111',parse_logic(canonical_value(row,'TYP',4),MODULE,SIGNAL,4))
        row['MAX']='2222';before=copy.deepcopy(row)
        with self.assertRaises(ValueError): change_radix(row,'十六进制',resolved,True)
        self.assertEqual(before,row)

    def test_old_mixed_literals_and_decimal_meaning(self):
        row=dict(signal='txzqcal_pd[7:0]',MAX="8'h4C",TYP="8'd76",MIN='0b01001100')
        display=display_row(row,self.data['resolved'],True)
        self.assertEqual('8',display['width'])
        self.assertEqual('十六进制',display['radix'])
        self.assertEqual(['4C']*3,[display[c] for c in CORNERS])
        old=dict(signal=SIGNAL,MAX='10',TYP="4'b1111",MIN='0xF')
        shown=display_row(old,self.data['resolved'],True)
        self.assertEqual('十进制',shown['radix'])
        self.assertEqual('10',shown['MAX'])
        change_radix(old,'二进制',self.data['resolved'],True)
        self.assertEqual('1010',old['MAX'])
        self.assertEqual('1111',old['TYP'])
        invalid=dict(signal=SIGNAL,MAX='1111',TYP='',MIN='')
        with self.assertRaises(ValueError): change_radix(invalid,'二进制',self.data['resolved'],True)

    def test_immediate_width_errors_and_quote_compatibility(self):
        for raw in ("4'b1111",'4’b1111','4＇b1111',"4 ' b1111",'0b1111','0xF',"4'd15"):
            self.assertEqual(15,number(raw,'二进制',4))
            self.assertEqual('1111',parse_logic(raw,MODULE,SIGNAL,4))
        for raw,base in [('1F','十六进制'),('2222','二进制'),('16','十进制'),('4“b1111','二进制')]:
            row=self.row(base,raw)
            with self.assertRaises(ValueError): number(raw,base,4)
            data=copy.deepcopy(self.data);data['signals']=[row]
            checked=precheck(data,TEMPLATES)
            self.assertTrue(any(r['code']=='SIGNAL_VALUE' for r in checked['issues']))
        row=dict(signal=SIGNAL,MAX='1111',TYP='',MIN='')
        edit_code(row,'MAX','15',self.data['resolved'],True)
        self.assertEqual('十进制',row['radix'])
        self.assertEqual("4'd15",canonical_value(row,'MAX',4))

    def test_excel_roundtrip_keeps_radix_values_and_width_from_msi(self):
        for radix,code in [('二进制','1111'),('十进制','15'),('十六进制','F')]:
            data=copy.deepcopy(self.data)
            data['signals']=[self.row(radix,code)]
            data['overrides']=[dict(self.row(radix,code),module=MODULE,direction='DRV')]
            exported=write(data,self.root/'codes.xlsx')
            book=load_workbook(exported)
            sheet=book['Config_Generated']
            self.assertTrue(any(c.value=='进制' for row in sheet for c in row))
            # A human cannot spoof MSI WIDTH by editing the displayed Excel cell.
            for row in sheet:
                if row[0].value=='Signal': row[6].value='100'
            book.save(exported);book.close()
            loaded=precheck(read(exported),TEMPLATES)
            self.assertReady(loaded)
            for field in ('signals','overrides'):
                row=next(r for r in loaded[field] if r.get('signal')==SIGNAL)
                self.assertEqual(radix,row['radix'])
                self.assertEqual(code,row['TYP'])
                self.assertEqual('4',display_row(row,loaded['resolved'],field=='signals')['width'])
                self.assertEqual('1111',parse_logic(canonical_value(row,'MAX',4),MODULE,SIGNAL,4))

    def test_excel_base_only_change_converts_unchanged_cells(self):
        data=copy.deepcopy(self.data)
        data['signals']=[self.row('二进制','10')]
        data['overrides']=[dict(self.row('二进制','10'),module=MODULE,direction='DRV')]
        path=write(data,self.root/'base-change.xlsx')
        book=load_workbook(path)
        for row in book['Config_Generated']:
            if row[0].value=='Signal': row[7].value='十进制'
            if row[0].value==MODULE and row[2].value==SIGNAL:
                row[7].value='十六进制'
                row[3].value='F'  # Explicitly edited MAX becomes 15; unchanged TYP stays 2.
        book.save(path);book.close()
        checked=precheck(read(path),TEMPLATES)
        self.assertReady(checked)
        self.assertEqual('2',checked['signals'][0]['TYP'])
        local=next(r for r in checked['overrides'] if r.get('signal')==SIGNAL)
        self.assertEqual('2',local['TYP'])
        self.assertEqual('F',local['MAX'])
        self.assertEqual('0010',parse_logic(canonical_value(local,'TYP',4),MODULE,SIGNAL,4))

    def test_input_new_columns_and_old_input_preserves_decimal(self):
        data=copy.deepcopy(self.data);data['signals']=[dict(signal=SIGNAL,MAX='10',TYP='10',MIN='10')]
        path=input_book(data,self.root/'input.xlsx')
        old=precheck(read(path),TEMPLATES)
        self.assertEqual('1010',parse_logic(core_config(old).overrides[0]['value'],MODULE,SIGNAL,4))
        book=load_workbook(path);sheet=book['Input']
        header=next(row[0].row for row in sheet if row[0].value=='Signal')
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row >= header:
                sheet.unmerge_cells(str(merged))
        sheet.cell(header,5,'位宽');sheet.cell(header,6,'进制')
        for row in sheet:
            if row[0].value==SIGNAL: sheet.cell(row[0].row,6,'二进制')
        book.save(path);book.close()
        new=precheck(read(path),TEMPLATES)
        self.assertEqual('0010',parse_logic(core_config(new).overrides[0]['value'],MODULE,SIGNAL,4))

    def test_msi_default_and_override_different_syntax_not_double_interpreted(self):
        path=Path(self.data['msi_mappings'][0]['file'])
        book=load_workbook(path)
        for row in book[MODULE]:
            if row[0].value==SIGNAL: row[7].value='4’b1111'
        book.save(path);book.close()
        # Reconfirm after a deliberate MSI change; no edit to frozen input.
        data=precheck(self.data,TEMPLATES)
        for field in ('pins','msi_mappings'):
            for row in data[field]: row.update(status='MANUAL',fingerprint=row['current_fingerprint'])
        data['overrides']=[dict(self.row('十六进制','F'),module=MODULE,direction='DRV')]
        checked=precheck(data,TEMPLATES)
        self.assertReady(checked)
        self.assertFalse(any(r['code'] in ('SIGNAL_VALUE','MSI_EFFECTIVE_VALUE') for r in checked['issues']))
        result=generate(ProjectStore(self.root/'workspace'),ResourceResolver.current(),ProjectDocument.create('双链路','ui',payload(checked)))
        self.assertTrue(result.ready)
        for case in result.cases:
            corner=next(Path(case.path).glob('*_typ.sp')).read_text()
            self.assertIn('vreg_txffe3 reg_txffe[3] 0 dc vdd_typ',corner)

    def test_append_uses_confirmed_effective_pins_per_direction(self):
        data=copy.deepcopy(self.data)
        data['appends']=[dict(module=MODULE,direction=d,body='Res ipadc vss r=100') for d in ('DRV','RCV')]
        checked=precheck(data,TEMPLATES)
        self.assertReady(checked)
        self.assertFalse(any(r['code']=='APPEND_SYNTAX' for r in checked['issues']))
        project=ProjectDocument.create('映射后追加','ui',payload(checked))
        result=generate(ProjectStore(self.root/'workspace'),ResourceResolver.current(),project)
        self.assertTrue(result.ready)
        from ibis_auto.spf import parse
        top,pins=parse(Path(checked['resolved']['spf'][MODULE]['files']['TYP']))
        for case in result.cases:
            text=(Path(case.path)/'ckt_topology.sp').read_text()
            self.assertIn('Res ipadc vss r=100',text)
            actual=' '.join(line[3:] if line.startswith('x1 ') else line[2:] for line in text.splitlines() if line.startswith(('x1 ','+ '))).split()
            direction=case.key.split('/')[1]
            aliases={r['spf_pin']:r['t2b_pin'] for r in checked['pins'] if r['direction']==direction}
            self.assertEqual([aliases.get(p,p) for p in pins]+[top],actual)
        for row in data['pins']:
            if row['direction']=='RCV' and row['t2b_pin']=='ipadc': row['status']='UNRESOLVED'
        checked=precheck(data,TEMPLATES)
        warnings=[r for r in checked['issues'] if r['code']=='APPEND_SYNTAX']
        self.assertEqual(['RCV'],[r['direction'] for r in warnings])
        data['appends'][0]['body']='Res nonexistent vss r=100'
        self.assertTrue(any('nonexistent' in r['message'] for r in precheck(data,TEMPLATES)['issues']))
        # Stale mappings must not be treated as confirmed aliases.
        for row in data['pins']:
            if row['t2b_pin']=='ipadc': row['status']='STALE'
        data['appends'][0]['body']='Res ipadc vss r=100'
        self.assertEqual({'DRV','RCV'},{r['direction'] for r in precheck(data,TEMPLATES)['issues'] if r['code']=='APPEND_SYNTAX'})


@unittest.skipUnless(os.name=='nt' or os.environ.get('DISPLAY'), 'Tk 窗口测试需要 DISPLAY')
class WindowsFixGuiTests(unittest.TestCase):
    def test_two_pane_preview_selection_copy_and_folder(self):
        import tkinter as tk
        from tkinter import font
        from ibis_ui.file_preview import CaseFilePreview
        with tempfile.TemporaryDirectory() as temporary:
            directory=Path(temporary)
            for name in ('z.ini','a.t2b','corner_drv_typ.sp','component.sp','ckt_topology.sp'):
                (directory/name).write_bytes(('  '+name+'\tvalue\r\n  preserved  \r\n').encode())
            root=tk.Tk();root.withdraw()
            window=CaseFilePreview(root,directory)
            try:
                root.update()
                self.assertEqual(['ckt_topology.sp','component.sp','corner_drv_typ.sp','a.t2b','z.ini'],[p.name for p in window.files])
                self.assertEqual('none',str(window.preview['wrap']))
                self.assertEqual('TkFixedFont',str(window.preview['font']))
                self.assertEqual((directory/'ckt_topology.sp').read_bytes().decode(),window.preview.get('1.0','end-1c'))
                window.preview.focus_force();root.update()
                window.preview.event_generate('<Control-a>');window.preview.event_generate('<Control-c>');root.update()
                self.assertEqual(window.preview.get('1.0','end-1c'),window.clipboard_get())
                window.listing.selection_clear(0,'end');window.listing.selection_set(1)
                window.listing.event_generate('<<ListboxSelect>>');root.update()
                self.assertIn('component.sp',window.preview.get('1.0','end'))
                with patch('ibis_ui.file_preview.open_directory') as opener:
                    window.open_folder();opener.assert_called_once_with(directory.resolve())
            finally: window.destroy();root.destroy()

    def test_cell_code_edit_switch_error_and_readonly_width(self):
        from ibis_ui.controller import ProjectController
        from ibis_ui.view import IBISApplicationView
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);data=confirmed(fixture(root))
            data['signals']=[dict(signal=SIGNAL,encoding='code',radix='二进制',MAX='',TYP='',MIN='')]
            controller=ProjectController(ProjectStore(root/'workspace'))
            controller._activate(ProjectDocument.create('界面','ui',payload(data)),unsaved=True)
            app=IBISApplicationView(controller)
            try:
                app._show_main();app.update()
                panel=app.vnext_panel;editor=panel.editors['signals']
                self.assertEqual('4',editor.rows()[0]['width'])
                editor.edit_cell(0,'width');self.assertIsNone(editor.editor)
                editor.set_value(0,'MAX','1111')
                editor.set_value(0,'radix','十六进制')
                self.assertEqual('F',editor.rows()[0]['MAX'])
                editor.set_value(0,'radix','十进制')
                self.assertEqual('15',editor.rows()[0]['MAX'])
                editor.set_value(0,'radix','二进制')
                self.assertEqual('1111',editor.rows()[0]['MAX'])
                panel.check();self.assertEqual('normal',str(panel.generate_button['state']))
                editor.set_value(0,'radix','十六进制');editor.set_value(0,'MAX','1F')
                self.assertIn('4-bit',str(editor.edit_error['text']))
                panel.check();self.assertEqual('disabled',str(panel.generate_button['state']))
                controller.export_config(root/'invalid.xlsx')
                editor.set_value(0,'MAX','F');panel.check()
                self.assertEqual('normal',str(panel.generate_button['state']))
                controller.save_configuration()
                self.assertEqual('1111',parse_logic(core_config(controller.active_config['vnext']).overrides[0]['value'],MODULE,SIGNAL,4))
            finally: app.destroy()
